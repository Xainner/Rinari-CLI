"""Documentos sintéticos para las pruebas: pequeños, en español y exigentes."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path


def deck(path: Path, *, slides: int = 2, chart: bool = True, notes: bool = True) -> Path:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    presentation = Presentation()
    for index in range(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = f"Diapositiva {index + 1}: crecimiento del año ñandú"
        box = slide.shapes.add_textbox(Inches(1), Inches(1.6), Inches(8), Inches(1))
        box.text_frame.text = "Texto del cuerpo con cifras: 12,5 % y \u22123"
        if chart and index == 0:
            data = CategoryChartData()
            data.categories = ["Norte", "Sur"]
            data.add_series("2026", (12.5, 8))
            slide.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(6), Inches(3), data
            )
        if notes:
            slide.notes_slide.notes_text_frame.text = f"Nota {index + 1}"
    presentation.save(path)
    return path


def pdf(path: Path, *, pages: int = 2) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    drawing = canvas.Canvas(str(path), pagesize=A4)
    for index in range(pages):
        drawing.drawString(72, 760, f"Página {index + 1} — informe")
        drawing.showPage()
    drawing.save()
    return path


def with_content_type(source: Path, target: Path, old: str, new: str) -> Path:
    """Copia de un paquete OOXML con otro tipo principal (p. ej. un .pptm)."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(buffer, "w") as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(old.encode(), new.encode())
            dst.writestr(info, data)
    target.write_bytes(buffer.getvalue())
    return target
