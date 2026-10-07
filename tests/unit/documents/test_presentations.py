"""Entrega 2: presentaciones de calidad de producción.

Autoría desde un DeckSpec, edición tipada con diff de preservación, revisión
visual ligada a los renders de la revisión exacta y finalización que no
aprueba lo que no se comprobó.
"""

from __future__ import annotations

import dataclasses
import io

import pytest

from rinari.artifacts.store import ArtifactStore
from rinari.documents import preservation
from rinari.documents.adapters import pptx_build, pptx_edit, render
from rinari.documents.contracts import DocumentError, DocumentErrorCode
from rinari.documents.design.tokens import THEMES, contrast
from rinari.documents.jobs import JobManager
from rinari.documents.service import DocumentService
from rinari.documents.validation import pptx_checks
from tests.unit import test_tool_runtime as base
from tests.unit.test_tool_runtime import _ctx, _runtime

project = base.project
SESSION = "ses_decks"

DECK = {
    "title": "Revisión comercial Q3",
    "language": "es",
    "theme": "executive-light",
    "slides": [
        {
            "layout": "cover",
            "title": "Resultados comerciales del tercer trimestre",
            "subtitle": "Crecimiento, márgenes y prioridades",
        },
        {
            "layout": "summary",
            "title": "La región norte explica dos tercios del crecimiento",
            "points": [
                {"head": "Ventas +18 %", "text": "El trimestre cerró en 4,2 M."},
                {"head": "Riesgo en Sur", "text": "Tres cuentas renegocian contrato."},
            ],
        },
        {
            "layout": "chart",
            "title": "Norte y Centro concentran el crecimiento",
            "chart": {
                "type": "column",
                "categories": ["Norte", "Centro", "Sur", "Oeste"],
                "series": [{"name": "Crecimiento %", "values": [31, 19, 4, -3]}],
                "number_format": '0"%"',
            },
            "takeaway": "Priorizar Norte y estabilizar Oeste.",
            "source": "Fuente: CRM, ventas cerradas.",
        },
        {
            "layout": "table",
            "title": "Detalle por región",
            "columns": ["Región", "Ventas (M)", "Margen %"],
            "rows": [["Norte", 1.62, "39,1 %"], ["Sur", 0.71, "37,2 %"]],
        },
        {"layout": "closing", "title": "Gracias"},
    ],
}


@pytest.fixture
def store(app_ctx, monkeypatch):
    # Sin renderer: lo visual queda «no ejecutado», que es lo que se prueba aquí.
    monkeypatch.setenv("RINARI_DOCUMENTS_NO_OFFICE", "1")
    monkeypatch.setenv("RINARI_DOCUMENTS_NO_LIBREOFFICE", "1")
    yield ArtifactStore(app_ctx)
    JobManager.close_for(app_ctx)


def _created(service: DocumentService, spec=DECK) -> dict:
    job = service.wait(service.create(spec), 120)
    assert job["status"] == "succeeded", job["error"]
    return job["result"]


def _fake_renders(store, service: DocumentService, revision_id: str, pages: int) -> None:
    """Capturas sintéticas de esta revisión, como las dejaría un render real."""
    from PIL import Image

    for page in range(1, pages + 1):
        buffer = io.BytesIO()
        Image.new("RGB", (64, 36), (255, 255, 255)).save(buffer, "PNG")
        store.create(
            SESSION,
            "previews",
            f"{revision_id}-pdfium-p{page:04d}.png",
            buffer.getvalue(),
            content_type="image/png",
        )


# -- diseño --------------------------------------------------------------------------


def test_themes_have_readable_text_and_semantic_colors():
    for theme in THEMES.values():
        assert contrast(theme.text, theme.background) >= 7, theme.id
        assert contrast(theme.muted, theme.background) >= 4.5, theme.id
        assert contrast(theme.accent_text, theme.accent) >= 4.5, theme.id
        assert contrast(theme.ink, theme.background) >= 4.5, theme.id
        for semantic in (theme.positive, theme.negative):
            assert contrast(semantic, theme.surface) >= 4.5, (theme.id, semantic)


@pytest.mark.parametrize("theme_id", sorted(THEMES))
def test_every_theme_builds_a_deck_without_static_findings(theme_id, tmp_path):
    data, plans = pptx_build.build({**DECK, "theme": theme_id})
    path = tmp_path / "deck.pptx"
    path.write_bytes(data)
    assert pptx_build.plan_findings(plans) == []
    assert pptx_checks.structure(path).status == "passed"
    layout = pptx_checks.layout(path)
    assert layout.status == "passed", layout.findings


def test_specs_are_closed_and_layouts_named(tmp_path):
    with pytest.raises(DocumentError) as err:
        pptx_build.build({**DECK, "colour": "red"})
    assert err.value.code is DocumentErrorCode.INVALID_SPEC
    bad = {"slides": [{"layout": "bullets", "title": "x", "points": ["a"], "subtitle": "y"}]}
    with pytest.raises(DocumentError) as err:
        pptx_build.build(bad)
    assert "subtitle" in err.value.message


def test_slide_numbers_are_fields_and_shapes_are_named(tmp_path):
    from pptx import Presentation

    data, _ = pptx_build.build(DECK)
    path = tmp_path / "deck.pptx"
    path.write_bytes(data)
    slide = Presentation(str(path)).slides[1]
    names = {shape.name for shape in slide.shapes}
    assert {"rinari:s02:title", "rinari:s02:number"} <= names
    number = next(s for s in slide.shapes if s.name == "rinari:s02:number")
    assert 'type="slidenum"' in number._element.xml


def test_long_text_is_reported_not_shrunk_to_illegible():
    long = " ".join(["palabra"] * 400)
    _, plans = pptx_build.build({"slides": [{"layout": "statement", "title": long}]})
    findings = pptx_build.plan_findings(plans)
    assert findings and findings[0]["code"] == "TEXT_OVERFLOW"


# -- edición -----------------------------------------------------------------------


@pytest.fixture
def deck(tmp_path):
    data, _ = pptx_build.build(DECK)
    path = tmp_path / "deck.pptx"
    path.write_bytes(data)
    return path


def test_replace_text_keeps_the_formatting_of_split_runs(tmp_path):
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, Pt(400), Pt(50))
    paragraph = box.text_frame.paragraphs[0]
    for text, bold in (("Ventas de No", False), ("rte", True), (" en alza", False)):
        run = paragraph.add_run()
        run.text = text
        run.font.bold = bold
    path = tmp_path / "runs.pptx"
    prs.save(path)

    pptx_edit.apply(
        path,
        tmp_path / "out.pptx",
        [{"op": "pptx.replace_text", "find": "Norte", "replace": "Sur"}],
        {},
    )
    out = Presentation(str(tmp_path / "out.pptx")).slides[0].shapes[0].text_frame.paragraphs[0]
    assert "".join(r.text for r in out.runs) == "Ventas de Sur en alza"
    assert out.runs[-1].text == " en alza" and not out.runs[-1].font.bold


def test_edits_touch_only_the_declared_parts(deck, tmp_path):
    from pptx import Presentation

    operations = [
        {
            "op": "pptx.set_text",
            "slide": 2,
            "shape_name": "rinari:s02:title",
            "text": "Norte sigue liderando",
            "expected_text": "La región norte explica dos tercios del crecimiento",
        },
        {
            "op": "pptx.update_chart",
            "slide": 3,
            "shape_name": "rinari:s03:chart",
            "categories": ["Norte", "Centro", "Sur", "Oeste"],
            "series": [{"name": "Crecimiento %", "values": [33, 18, 5, -2]}],
        },
        {
            "op": "pptx.set_table_cell",
            "slide": 4,
            "shape_name": "rinari:s04:table",
            "row": 1,
            "col": 1,
            "text": "1,70",
            "expected_text": "1,62",
        },
    ]
    out = tmp_path / "out.pptx"
    state = pptx_edit.apply(deck, out, operations, {})
    diff = preservation.diff_packages(deck, out)
    verdict = preservation.evaluate(diff, state.allowed, "preserve_strict")
    assert verdict["status"] == "passed", verdict
    changed = {row["part"] for row in diff["changed"]}
    assert "ppt/slides/slide2.xml" in changed and "ppt/slides/slide1.xml" not in changed
    chart = Presentation(str(out)).slides[2].shapes
    plot = next(s for s in chart if s.has_chart).chart.plots[0]
    assert list(plot.series[0].values) == [33, 18, 5, -2]
    # El formato de las cifras del gráfico sigue siendo el original.
    assert '0"%"' in next(s for s in chart if s.has_chart).chart._chartSpace.xml


def test_a_stale_precondition_is_a_conflict(deck, tmp_path):
    with pytest.raises(DocumentError) as err:
        pptx_edit.apply(
            deck,
            tmp_path / "out.pptx",
            [
                {
                    "op": "pptx.set_text",
                    "slide": 2,
                    "shape_name": "rinari:s02:title",
                    "text": "x",
                    "expected_text": "otro título",
                }
            ],
            {},
        )
    assert err.value.code is DocumentErrorCode.REVISION_CONFLICT
    assert not (tmp_path / "out.pptx").exists()


def test_slides_can_be_added_moved_and_deleted(deck, tmp_path):
    from pptx import Presentation

    out = tmp_path / "out.pptx"
    operations = [
        {
            "op": "pptx.add_slide",
            "after": 1,
            "spec": {"layout": "bullets", "title": "Agenda", "points": ["Uno", "Dos"]},
        },
        {
            "op": "pptx.delete_slide",
            "slide": 4,
            "expected_title": "Norte y Centro concentran el crecimiento",
        },
        {"op": "pptx.move_slide", "slide": 5, "to": 2},
    ]
    state = pptx_edit.apply(deck, out, operations, {})
    titles = [
        s.shapes.title.text
        if s.shapes.title
        else next((sh.text_frame.text for sh in s.shapes if sh.name.endswith(":title")), "")
        for s in Presentation(str(out)).slides
    ]
    assert titles[1] == "Gracias" and titles[2] == "Agenda"
    assert "Norte y Centro concentran el crecimiento" not in titles
    diff = preservation.diff_packages(deck, out)
    verdict = preservation.evaluate(
        diff, state.allowed, "preserve_strict", allow_new=tuple(state.allow_new)
    )
    assert verdict["status"] == "passed", verdict
    # El gráfico de la diapositiva borrada se va con ella.
    assert any(row["category"] == "chart" for row in diff["removed"])


def test_undeclared_changes_fail_strict_preservation():
    diff = {
        "added": [],
        "removed": [],
        "changed": [
            {"part": "ppt/slides/slide2.xml", "category": "slide"},
            {"part": "ppt/slideMasters/slideMaster1.xml", "category": "master"},
        ],
        "unchanged": 40,
    }
    strict = preservation.evaluate(diff, {"ppt/slides/slide2.xml"}, "preserve_strict")
    assert strict["status"] == "failed" and strict["risky"]
    assert preservation.evaluate(diff, {"ppt/slides/slide2.xml"}, "rebuild")["status"] == "passed"


def test_unknown_operation_fields_are_rejected():
    with pytest.raises(DocumentError) as err:
        pptx_edit.validate([{"op": "pptx.set_text", "slide": 1, "shape_id": 2, "txt": "x"}])
    assert err.value.code is DocumentErrorCode.INVALID_SPEC
    with pytest.raises(DocumentError):
        pptx_edit.validate([{"op": "pptx.run_macro"}])


# -- servicio: crear, editar, revisar, finalizar ---------------------------------


def test_create_is_a_draft_whose_visual_check_is_not_claimed(store):
    service = DocumentService(store, SESSION)
    result = _created(service)
    revision = result["revision"]
    assert revision["operation"] == "create" and revision["state"] == "draft"
    assert revision["spec_uri"].startswith(f"artifact://{SESSION}/documents/spec-")
    report = result["report"]
    assert report["checks"]["structure"]["status"] == "passed"
    assert report["checks"]["layout"]["status"] == "passed"
    assert report["checks"]["visual"] == {"status": "not_run", "reason": "RENDER_UNAVAILABLE"}
    assert result["render"]["status"] == "unavailable"
    assert report["status"] == "partial"


def test_finalize_blocks_until_every_page_was_reviewed(store):
    service = DocumentService(store, SESSION)
    revision = _created(service)["revision"]["id"]
    blocked = service.finalize(revision)
    assert blocked["finalized"] is False
    assert {"check": "visual", "status": "not_run", "reason": "RENDER_UNAVAILABLE"} in blocked[
        "pending"
    ]

    _fake_renders(store, service, revision, 5)
    with pytest.raises(DocumentError):
        service.review(revision, pages=[6])
    partial = service.review(revision, pages=[1, 2])
    assert partial["checks"]["visual"]["status"] == "partial"
    done = service.review(revision, pages=[3, 4, 5])
    assert done["checks"]["visual"]["status"] == "passed"

    final = service.finalize(revision)
    assert final["finalized"] is True and final["state"] == "final"
    assert final["report"]["deliverable_state"] == "final"


def test_a_blocking_visual_finding_fails_the_check(store):
    service = DocumentService(store, SESSION)
    revision = _created(service)["revision"]["id"]
    _fake_renders(store, service, revision, 5)
    report = service.review(
        revision,
        pages=[1, 2, 3, 4, 5],
        findings=[{"severity": "error", "page": 3, "message": "Etiqueta cortada", "fix": "..."}],
    )
    assert report["checks"]["visual"]["status"] == "failed"
    assert service.finalize(revision, accept_partial=True)["finalized"] is False


def test_accepting_a_partial_draft_keeps_the_label(store):
    service = DocumentService(store, SESSION)
    revision = _created(service)["revision"]["id"]
    accepted = service.finalize(revision, accept_partial=True)
    assert accepted["finalized"] is True and accepted["state"] == "accepted_draft"
    assert accepted["pending"]


def test_edit_makes_a_child_revision_and_reports_the_change(store):
    service = DocumentService(store, SESSION)
    parent = _created(service)["revision"]
    job = service.wait(
        service.edit(
            parent["id"],
            [{"op": "pptx.replace_text", "find": "Norte", "replace": "Septentrión"}],
            expected_sha256=parent["sha256"],
        ),
        120,
    )
    assert job["status"] == "succeeded", job["error"]
    child = job["result"]["revision"]
    assert child["parent_id"] == parent["id"] and child["document_id"] == parent["document_id"]
    report = job["result"]["report"]
    assert report["checks"]["preservation"]["status"] == "passed"
    assert report["changes"][0]["count"] >= 2
    assert any(row["change"] == "content_changed" for row in report["semantic_diff"])
    # El original sigue igual.
    assert service.revisions.get(parent["id"]).sha256 == parent["sha256"]
    diff = service.diff(parent["id"], child["id"])
    assert diff["content"] and diff["package"]["changed"]


def test_edit_with_a_stale_hash_is_refused_before_starting(store):
    service = DocumentService(store, SESSION)
    parent = _created(service)["revision"]
    with pytest.raises(DocumentError) as err:
        service.edit(
            parent["id"],
            [{"op": "pptx.set_notes", "slide": 1, "text": "x"}],
            expected_sha256="0" * 64,
        )
    assert err.value.code is DocumentErrorCode.REVISION_CONFLICT


def test_a_failed_edit_leaves_no_revision(store):
    service = DocumentService(store, SESSION)
    parent = _created(service)["revision"]
    before = len(service.list_revisions())
    job = service.wait(
        service.edit(
            parent["id"], [{"op": "pptx.replace_text", "find": "inexistente", "replace": "x"}]
        ),
        120,
    )
    assert job["status"] == "failed"
    assert job["error"]["code"] == "REVISION_CONFLICT"
    assert len(service.list_revisions()) == before


# -- herramientas --------------------------------------------------------------------


def _tool_runtime(app_ctx, tmp_path):
    store = ArtifactStore(app_ctx)
    root = tmp_path / "proj"
    root.mkdir(exist_ok=True)
    ctx = dataclasses.replace(
        _ctx(tmp_path, root), artifact_root=store._root(), artifact_store=store
    )
    runtime, _ = _runtime(ctx, tmp_path, answer="y")
    return store, root, ctx, runtime


def test_tools_create_edit_and_save_without_overwriting(store, app_ctx, tmp_path):
    _, root, ctx, runtime = _tool_runtime(app_ctx, tmp_path)
    catalog = runtime.execute("documents.templates", {"kind": "pptx"}, ctx)
    assert catalog.ok and "kpi" in catalog.data["pptx"]["layouts"]

    created = runtime.execute("documents.create", {"spec": DECK, "wait_s": 120}, ctx)
    assert created.ok, created.error
    revision = created.data["result"]["revision"]["id"]

    edited = runtime.execute(
        "documents.edit",
        {
            "document": revision,
            "operations": [{"op": "pptx.set_notes", "slide": 2, "text": "Insistir en volumen"}],
            "wait_s": 120,
        },
        ctx,
    )
    assert edited.ok and edited.data["status"] == "succeeded", edited.data
    child = edited.data["result"]["revision"]["id"]
    diff = runtime.execute("documents.diff", {"after": child}, ctx)
    assert diff.ok and diff.data["before"] == revision

    pending = runtime.execute("documents.finalize", {"document": child, "save_to": "out"}, ctx)
    assert pending.ok and pending.data["finalized"] is False
    assert not (root / "out").exists()

    (root / "Revision.pptx").write_bytes(b"del usuario")
    saved = runtime.execute(
        "documents.finalize",
        {"document": child, "accept_partial": True, "save_to": "Revision.pptx"},
        ctx,
    )
    assert saved.ok, saved.error
    assert (root / "Revision.pptx").read_bytes() == b"del usuario"
    assert saved.data["saved_to"].endswith("Revision (1).pptx")


def test_tools_reject_resources_outside_the_sandbox(store, app_ctx, tmp_path):
    _, _, ctx, runtime = _tool_runtime(app_ctx, tmp_path)
    outside = tmp_path / "fuera.png"
    outside.write_bytes(b"\x89PNG")
    spec = {"slides": [{"layout": "image", "title": "Foto", "image": "a"}]}
    result = runtime.execute(
        "documents.create", {"spec": spec, "resources": {"a": str(outside)}}, ctx
    )
    assert not result.ok
    assert result.error.message.startswith("UNSAFE_EXTERNAL_RESOURCE")


@pytest.mark.skipif(not render.office_renderers("pptx"), reason="no Office or LibreOffice")
def test_created_decks_render_with_the_installed_office(app_ctx, tmp_path):
    store = ArtifactStore(app_ctx)
    try:
        service = DocumentService(store, SESSION)
        job = service.wait(service.create(DECK), 300)
        assert job["status"] == "succeeded", job["error"]
        assert job["result"]["render"]["page_count"] == 5
        assert job["result"]["report"]["checks"]["visual"]["reason"] == "NOT_REVIEWED"
    finally:
        JobManager.close_for(app_ctx)
