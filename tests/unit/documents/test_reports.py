"""Entrega 4: Word y PDF profesionales.

Un ReportSpec produce el DOCX editable y el PDF de lectura; la edición de
Word conserva el formato de los runs y los comentarios; la manipulación de PDF
rellena formularios que se ven y no aplana nada sin pedirlo.
"""

from __future__ import annotations

import io

import pytest

from rinari.artifacts.store import ArtifactStore
from rinari.documents import preservation
from rinari.documents.adapters import (
    docx_build,
    docx_edit,
    pdf_build,
    pdf_edit,
    report_spec,
    word_fields,
)
from rinari.documents.contracts import DocumentError, DocumentErrorCode
from rinari.documents.jobs import JobManager
from rinari.documents.service import DocumentService
from rinari.documents.validation import report_checks

SESSION = "ses_reports"

REPORT = {
    "title": "Informe anual",
    "subtitle": "Resultados y recomendaciones",
    "author": "Dirección",
    "language": "es",
    "toc": True,
    "footer": "Confidencial",
    "blocks": [
        {"type": "heading", "text": "1. Resumen", "level": 1},
        {"type": "paragraph", "text": "Las ventas crecieron un **12 %** con margen *estable*."},
        {
            "type": "kpis",
            "items": [
                {"label": "Ventas", "value": "4,2 M", "delta": "+12 %"},
                {"label": "Margen", "value": "38 %", "delta": "\u22120,5 pp"},
            ],
        },
        {"type": "bullets", "items": ["Ampliar Norte", "Blindar Sur"]},
        {
            "type": "table",
            "caption": "Ventas por región",
            "columns": ["Región", "Ventas"],
            "rows": [["Norte", 1.62], ["Sur", 0.71]],
            "source": "ERP",
        },
        {
            "type": "callout",
            "title": "Decisión",
            "text": "Aprobar antes del 15 de diciembre.",
            "tone": "warning",
        },
        {"type": "heading", "text": "1.1 Detalle", "level": 2},
        {"type": "quote", "text": "El norte sigue tirando.", "author": "Dirección regional"},
        {"type": "appendix", "title": "Metodología"},
        {"type": "paragraph", "text": "Datos conciliados del ERP."},
    ],
}


@pytest.fixture
def store(app_ctx, monkeypatch):
    monkeypatch.setenv("RINARI_DOCUMENTS_NO_OFFICE", "1")
    monkeypatch.setenv("RINARI_DOCUMENTS_NO_LIBREOFFICE", "1")
    yield ArtifactStore(app_ctx)
    JobManager.close_for(app_ctx)


# -- autoría --------------------------------------------------------------------------


def test_report_specs_are_closed():
    with pytest.raises(DocumentError):
        report_spec.validate({**REPORT, "colour": "red"})
    with pytest.raises(DocumentError) as err:
        report_spec.validate({"blocks": [{"type": "heading", "text": "x", "size": 3}]})
    assert "size" in err.value.message


def test_docx_uses_real_styles_fields_and_repeated_headers(tmp_path):
    from docx import Document

    data, plan = docx_build.build(REPORT)
    path = tmp_path / "informe.docx"
    path.write_bytes(data)
    document = Document(str(path))
    styles = {p.style.name for p in document.paragraphs}
    assert {"Title", "Heading 1", "Heading 2", "List Bullet", "Caption", "Quote"} <= styles
    xml = document.element.xml
    assert " TOC " in xml and "SEQ Tabla" in xml and "w:tblHeader" in xml
    footer = document.sections[0].footer._element.xml
    assert " PAGE " in footer
    assert plan.toc and plan.tables == 1
    checks = report_checks.docx_checks(path, ["Ampliar Norte", "Metodología"])
    assert checks["structure"]["status"] == "passed"
    assert checks["content"]["status"] == "passed"
    assert checks["fields"] == {
        "status": "not_run",
        "reason": "FIELDS_NOT_UPDATED",
        "evidence": {"toc": 1},
    }


def test_pdf_has_selectable_text_embedded_fonts_and_a_paginated_toc(tmp_path):
    from pypdf import PdfReader

    data, plan = pdf_build.build(REPORT)
    path = tmp_path / "informe.pdf"
    path.write_bytes(data)
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "Ventas por región" in text and "\u22120,5 pp" in text
    # El índice tiene números de página: la segunda pasada los resolvió.
    toc_page = reader.pages[1].extract_text()
    assert "1. Resumen" in toc_page and "Anexo A. Metodología" in toc_page
    checks = report_checks.pdf_checks(path, ["Blindar Sur"])
    assert checks["structure"]["status"] == "passed"
    assert checks["text_layer"]["status"] == "passed", checks["text_layer"]
    assert checks["content"]["status"] == "passed"
    assert plan.pages >= 4 and plan.toc


# -- edición de Word -----------------------------------------------------------------


@pytest.fixture
def informe(tmp_path):
    from docx import Document

    document = Document()
    paragraph = document.add_paragraph("El total de ")
    bold = paragraph.add_run("Norte")
    bold.bold = True
    paragraph.add_run(" fue de 1,6 M.")
    document.add_paragraph("Segundo párrafo intacto.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Región"
    table.cell(1, 0).text = "Norte"
    table.cell(1, 1).text = "1,6"
    document.add_comment(document.paragraphs[1].runs, text="Revisar", author="Ana")
    path = tmp_path / "informe.docx"
    document.save(str(path))
    return path


def test_word_edits_keep_run_formatting_and_comments(informe, tmp_path):
    from docx import Document

    out = tmp_path / "out.docx"
    state = docx_edit.apply(
        informe,
        out,
        [
            {"op": "docx.replace_text", "find": "1,6 M", "replace": "1,7 M", "expected_count": 1},
            {
                "op": "docx.set_cell",
                "block": 3,
                "row": 1,
                "col": 1,
                "text": "1,7",
                "expected_text": "1,6",
            },
            {"op": "docx.insert_paragraph", "after": 2, "text": "Nuevo párrafo."},
        ],
    )
    document = Document(str(out))
    first = document.paragraphs[0]
    assert first.text == "El total de Norte fue de 1,7 M."
    assert [r.bold for r in first.runs] == [None, True, None]
    assert document.paragraphs[2].text == "Nuevo párrafo."
    assert len(document.comments) == 1
    diff = preservation.diff_packages(informe, out)
    verdict = preservation.evaluate(diff, state.allowed, "preserve_strict")
    assert verdict["status"] == "passed", verdict
    changes = report_checks.docx_semantic_diff(informe, out)
    assert any(c["change"] == "block_added" for c in changes)


def test_word_edit_preconditions(informe, tmp_path):
    with pytest.raises(DocumentError) as err:
        docx_edit.apply(
            informe,
            tmp_path / "o.docx",
            [{"op": "docx.set_paragraph", "block": 1, "text": "x", "expected_text": "otro"}],
        )
    assert err.value.code is DocumentErrorCode.REVISION_CONFLICT
    with pytest.raises(DocumentError):
        docx_edit.validate([{"op": "docx.set_paragraph", "block": 1, "txt": "x"}])


def test_templates_are_sandboxed_and_strict(tmp_path):
    from docx import Document

    template = tmp_path / "plantilla.docx"
    document = Document()
    document.add_paragraph("Estimado {{ cliente }}, total {{ total }}.")
    document.save(str(template))
    out = tmp_path / "carta.docx"
    variables = docx_edit.render_template(
        template, out, {"cliente": "<ACME & Co>", "total": "1.200 €"}
    )
    assert variables == ["cliente", "total"]
    assert Document(str(out)).paragraphs[0].text == "Estimado <ACME & Co>, total 1.200 €."
    with pytest.raises(DocumentError) as err:
        docx_edit.render_template(template, out, {"cliente": "x"})
    assert err.value.code is DocumentErrorCode.INVALID_SPEC
    evil = tmp_path / "evil.docx"
    document = Document()
    document.add_paragraph("{{ cliente.__class__.__mro__[1].__subclasses__() }}")
    document.save(str(evil))
    with pytest.raises(DocumentError):
        docx_edit.render_template(evil, out, {"cliente": "x"})


# -- PDF -------------------------------------------------------------------------------


@pytest.fixture
def form_pdf(tmp_path):
    from reportlab.pdfgen import canvas

    path = tmp_path / "solicitud.pdf"
    page = canvas.Canvas(str(path))
    page.drawString(72, 760, "Solicitud")
    page.acroForm.textfield(name="nombre", x=72, y=700, width=300, height=20)
    page.acroForm.choice(
        name="tipo", options=["Alta", "Baja"], value="Alta", x=72, y=650, width=120, height=20
    )
    page.showPage()
    page.drawString(72, 760, "Segunda página")
    page.showPage()
    page.save()
    return path


def test_pdf_forms_are_filled_visibly_and_stay_editable(form_pdf, tmp_path):
    out = tmp_path / "out.pdf"
    state = pdf_edit.apply(
        form_pdf,
        out,
        [{"op": "pdf.fill_form", "fields": {"nombre": "José Rojas", "tipo": "Baja"}}],
        {},
    )
    check = report_checks.pdf_forms(out, state.filled)
    assert check.status == "passed", check.findings
    from pypdf import PdfReader

    assert PdfReader(str(out)).get_fields()["nombre"].get("/V") == "José Rojas"
    assert state.flattened is False
    with pytest.raises(DocumentError) as err:
        pdf_edit.apply(
            form_pdf, tmp_path / "x.pdf", [{"op": "pdf.fill_form", "fields": {"apellido": "x"}}], {}
        )
    assert err.value.code is DocumentErrorCode.NOT_FOUND
    with pytest.raises(DocumentError):
        pdf_edit.apply(
            form_pdf, tmp_path / "x.pdf", [{"op": "pdf.fill_form", "fields": {"tipo": "Otro"}}], {}
        )


def test_pdf_pages_are_manipulated_without_touching_content(form_pdf, tmp_path):
    from pypdf import PdfReader

    out = tmp_path / "out.pdf"
    other = tmp_path / "anexo.pdf"
    data, _ = pdf_build.build(
        {"cover": False, "blocks": [{"type": "paragraph", "text": "Anexo adjunto"}]}
    )
    other.write_bytes(data)
    pdf_edit.apply(
        form_pdf,
        out,
        [
            {"op": "pdf.merge", "documents": ["anexo"], "position": 2},
            {"op": "pdf.rotate", "pages": "1", "degrees": 90},
            {"op": "pdf.set_metadata", "title": "Solicitud firmada"},
        ],
        {"anexo": str(other)},
    )
    reader = PdfReader(str(out))
    assert len(reader.pages) == 3 and "Anexo adjunto" in reader.pages[2].extract_text()
    assert reader.metadata.title == "Solicitud firmada"
    with pytest.raises(DocumentError):
        pdf_edit.apply(
            form_pdf, tmp_path / "x.pdf", [{"op": "pdf.delete_pages", "pages": "1-2"}], {}
        )


# -- servicio --------------------------------------------------------------------------


def test_service_creates_word_and_pdf_from_one_spec(store):
    service = DocumentService(store, SESSION)
    word = service.wait(service.create({**REPORT, "kind": "docx"}), 120)
    assert word["status"] == "succeeded", word["error"]
    report = word["result"]["report"]
    assert report["checks"]["fields"]["reason"] == "FIELDS_NOT_UPDATED"
    pending = service.finalize(word["result"]["revision"]["id"], accept_partial=False)
    assert any(item["check"] == "fields" for item in pending["pending"])

    pdf = service.wait(service.create({**REPORT, "kind": "pdf"}), 120)
    assert pdf["status"] == "succeeded", pdf["error"]
    assert pdf["result"]["revision"]["backend"] == "reportlab"
    # Un PDF siempre se puede renderizar (pdfium): la revisión visual es posible.
    assert pdf["result"]["render"]["page_count"] >= 4
    assert pdf["result"]["report"]["checks"]["text_layer"]["status"] == "passed"


def test_service_fills_a_form_and_validation_remembers_it(store, form_pdf):
    service = DocumentService(store, SESSION)
    uri = store.create(
        SESSION, "media", "solicitud.pdf", form_pdf.read_bytes(), content_type="application/pdf"
    ).uri()
    job = service.wait(
        service.edit(uri, [{"op": "pdf.fill_form", "fields": {"nombre": "Ana"}}]), 120
    )
    assert job["status"] == "succeeded", job["error"]
    assert job["result"]["report"]["checks"]["forms"]["status"] == "passed"
    again = service.validate(job["result"]["revision"]["id"])
    assert again["checks"]["forms"]["status"] == "passed"


@pytest.mark.skipif(not word_fields.available(), reason="no Word")
def test_word_paginates_the_table_of_contents(app_ctx):
    store = ArtifactStore(app_ctx)
    try:
        service = DocumentService(store, SESSION)
        job = service.wait(service.create({**REPORT, "kind": "docx"}, render=False), 300)
        assert job["status"] == "succeeded", job["error"]
        assert job["result"]["report"]["checks"]["fields"]["status"] == "passed"
        assert job["result"]["revision"]["backend"] == "python-docx+word-com"
    finally:
        JobManager.close_for(app_ctx)
