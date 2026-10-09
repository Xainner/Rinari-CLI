"""Entrega 1: contenedores seguros, revisiones inmutables, PPTX, trabajos y render."""

from __future__ import annotations

import io
import threading
import time
import zipfile

import pytest

from rinari.artifacts.store import ArtifactStore
from rinari.documents import capabilities
from rinari.documents.adapters import render
from rinari.documents.contracts import DocumentError, DocumentErrorCode
from rinari.documents.jobs import JobManager
from rinari.documents.security import inspect_container, sniff
from rinari.documents.service import DocumentService
from tests.unit.documents import builders

PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
SESSION = "ses_docs"


@pytest.fixture
def store(app_ctx):
    yield ArtifactStore(app_ctx)
    # Ningún hilo de trabajo debe seguir usando la base al cerrarla.
    JobManager.close_for(app_ctx)


def _artifact(store, path, name=None):
    return store.create(
        SESSION, "media", name or path.name, path.read_bytes(), content_type=PPTX
    ).uri()


# -- contenedores ------------------------------------------------------------------


def test_format_comes_from_the_content_not_the_extension(tmp_path):
    source = builders.deck(tmp_path / "deck.pptx")
    macro = builders.with_content_type(
        source,
        tmp_path / "renamed.pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
        "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
    )
    assert sniff(source) == "pptx"
    assert sniff(macro) == "pptm"
    report = inspect_container(macro)
    assert report.has("macros")


def test_encrypted_and_legacy_office_files_are_named(tmp_path):
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600
    legacy = tmp_path / "viejo.ppt"
    legacy.write_bytes(ole)
    encrypted = tmp_path / "cifrado.pptx"
    encrypted.write_bytes(ole + "EncryptionInfo".encode("utf-16-le"))
    with pytest.raises(DocumentError) as old:
        sniff(legacy)
    with pytest.raises(DocumentError) as locked:
        sniff(encrypted)
    assert old.value.code is DocumentErrorCode.UNSUPPORTED_FORMAT
    assert locked.value.code is DocumentErrorCode.ENCRYPTED_INPUT


def test_unsafe_zip_paths_are_rejected(tmp_path):
    source = builders.deck(tmp_path / "deck.pptx")
    buffer = io.BytesIO()
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(buffer, "w") as dst:
        for info in src.infolist():
            dst.writestr(info, src.read(info.filename))
        dst.writestr("../escape.txt", "x")
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(buffer.getvalue())
    with pytest.raises(DocumentError) as err:
        inspect_container(bad)
    assert err.value.code is DocumentErrorCode.UNSAFE_EXTERNAL_RESOURCE


def test_native_chart_data_is_not_reported_as_an_ole_object(tmp_path):
    report = inspect_container(builders.deck(tmp_path / "deck.pptx", chart=True))
    assert not report.has("embedded_objects")


# -- revisiones y lectura ----------------------------------------------------------


def test_import_is_idempotent_and_tampering_is_caught(store, tmp_path):
    uri = _artifact(store, builders.deck(tmp_path / "deck.pptx"))
    service = DocumentService(store, SESSION)
    first = service.resolve(uri)
    again = service.resolve(uri)
    assert first.id == again.id and first.operation == "import" and first.kind == "pptx"
    stored = store._storage_path(store.meta(uri).storage_path)
    stored.write_bytes(b"changed")
    with pytest.raises(DocumentError) as err:
        service.inspect(first.id)
    assert err.value.code is DocumentErrorCode.REVISION_CONFLICT


def test_other_sessions_cannot_resolve_a_revision(store, tmp_path):
    uri = _artifact(store, builders.deck(tmp_path / "deck.pptx"))
    revision = DocumentService(store, SESSION).resolve(uri)
    with pytest.raises(DocumentError) as err:
        DocumentService(store, "ses_other").resolve(revision.id)
    assert err.value.code is DocumentErrorCode.NOT_FOUND


def test_pptx_inspection_and_reading_with_stable_ids(store, tmp_path):
    uri = _artifact(store, builders.deck(tmp_path / "deck.pptx", slides=3))
    service = DocumentService(store, SESSION)
    inspected = service.inspect(uri)["inspection"]
    assert inspected["slide_count"] == 3
    assert inspected["features"]["charts"] == 1 and inspected["features"]["notes"] == 3
    assert inspected["slides"][0]["title"].startswith("Diapositiva 1")
    read = service.read(uri, "1")
    shapes = read["slides"][0]["shapes"]
    chart = next(s for s in shapes if s["kind"] == "chart")
    assert chart["chart"]["categories"] == ["Norte", "Sur"]
    assert chart["chart"]["series"][0]["values"] == [12.5, 8.0]
    assert all(isinstance(s["shape_id"], int) for s in shapes)
    assert read["slides"][0]["notes"] == "Nota 1"


def test_reading_past_the_budget_returns_a_cursor(store, tmp_path, monkeypatch):
    from rinari.documents.adapters import pptx_read

    monkeypatch.setattr(pptx_read, "READ_CHAR_BUDGET", 500)
    uri = _artifact(store, builders.deck(tmp_path / "deck.pptx", slides=4, chart=False))
    service = DocumentService(store, SESSION)
    first = service.read(uri)
    assert first["truncated"] and first["next_cursor"] == 2
    second = service.read(uri, cursor=first["next_cursor"])
    assert second["slides"][0]["index"] == 2


def test_selection_outside_the_deck_is_an_error(store, tmp_path):
    uri = _artifact(store, builders.deck(tmp_path / "deck.pptx", slides=2))
    with pytest.raises(DocumentError) as err:
        DocumentService(store, SESSION).read(uri, "3")
    assert err.value.code is DocumentErrorCode.INVALID_SPEC


# -- trabajos --------------------------------------------------------------------


def test_jobs_report_phases_and_can_be_cancelled(store):
    service = DocumentService(store, SESSION)
    started = threading.Event()

    def runner(handle):
        handle.phase("build", done=0, total=10)
        started.set()
        for step in range(200):
            handle.check()
            handle.progress(step % 10, 10)
            time.sleep(0.02)
        return {"done": True}

    job = service.jobs.start(session_id=SESSION, operation="test", request={}, runner=runner)
    assert started.wait(5)
    service.cancel(job["job_id"])
    final = service.wait(job, 5)
    assert final["status"] == "cancelled"


def test_active_jobs_become_interrupted_after_a_restart(store, app_ctx):
    service = DocumentService(store, SESSION)
    release = threading.Event()
    job = service.jobs.start(
        session_id=SESSION, operation="test", request={}, runner=lambda h: release.wait(5) or {}
    )
    JobManager._instances.clear()  # otro proceso del Engine
    restarted = JobManager.for_context(app_ctx)
    assert restarted.get(job["job_id"], session_id=SESSION)["status"] == "interrupted"
    release.set()


def test_worker_errors_keep_their_code(store):
    service = DocumentService(store, SESSION)

    def runner(handle):
        return handle.run_worker("does-not-exist", {})

    job = service.wait(
        service.jobs.start(session_id=SESSION, operation="test", request={}, runner=runner), 60
    )
    assert job["status"] == "failed"
    assert job["error"]["code"] == "UNSUPPORTED_FEATURE"


# -- render ----------------------------------------------------------------------


def test_pdf_renders_with_pdfium_and_is_cached_per_revision(store, tmp_path):
    path = builders.pdf(tmp_path / "informe.pdf", pages=3)
    uri = store.create(SESSION, "media", "informe.pdf", path.read_bytes()).uri()
    service = DocumentService(store, SESSION)
    job = service.wait(service.render(uri), 120)
    assert job["status"] == "succeeded", job["error"]
    result = job["result"]
    assert result["backend"] == "native" and result["page_count"] == 3
    assert [p["page"] for p in result["pages"]] == [1, 2, 3]
    assert store.get(result["pages"][0]["uri"])[:8] == b"\x89PNG\r\n\x1a\n"
    again = service.wait(service.render(uri), 30)
    assert again["result"]["cached"] is True


def test_without_an_office_renderer_the_preview_is_declared_unavailable(
    store, tmp_path, monkeypatch
):
    monkeypatch.setenv("RINARI_DOCUMENTS_NO_OFFICE", "1")
    monkeypatch.setenv("RINARI_DOCUMENTS_NO_LIBREOFFICE", "1")
    uri = _artifact(store, builders.deck(tmp_path / "deck.pptx"))
    row = capabilities.capability("pptx", "render")
    assert row["available"] is False and row["reason"] == "BACKEND_UNAVAILABLE"
    with pytest.raises(DocumentError) as err:
        DocumentService(store, SESSION).render(uri)
    assert err.value.code is DocumentErrorCode.BACKEND_UNAVAILABLE


@pytest.mark.skipif(not render.office_renderers("pptx"), reason="no Office or LibreOffice")
def test_pptx_renders_through_the_installed_office(store, tmp_path):
    uri = _artifact(store, builders.deck(tmp_path / "deck.pptx", slides=2))
    service = DocumentService(store, SESSION)
    job = service.wait(service.render(uri), 300)
    assert job["status"] == "succeeded", job["error"]
    assert job["result"]["page_count"] == 2
    assert job["result"]["backend"] in ("office-com", "libreoffice")


def test_capabilities_never_claim_what_is_missing():
    rows = {(r["kind"], r["operation"]): r for r in capabilities.capabilities()}
    assert rows[("pptx", "inspect")]["available"] is True
    assert rows[("pdf", "redact")]["backend"] == "raster-pdfium"
    assert rows[("docx", "redact")]["available"] is False
    assert rows[("docx", "calculate")]["available"] is False


def test_pdf_docx_and_xlsx_are_inspected_and_read_in_bounds(store, tmp_path):
    from docx import Document
    from openpyxl import Workbook

    service = DocumentService(store, SESSION)
    pdf = builders.pdf(tmp_path / "a.pdf", pages=3)
    pdf_uri = store.create(SESSION, "media", "a.pdf", pdf.read_bytes()).uri()
    assert service.inspect(pdf_uri)["inspection"]["page_count"] == 3
    page = service.read(pdf_uri, "2")["pages"][0]
    assert page["page"] == 2 and page["has_text"] and "informe" in page["text"]

    document = Document()
    document.add_heading("Informe anual", 1)
    document.add_paragraph("Primer párrafo")
    document.add_table(2, 2).cell(0, 0).text = "celda"
    document.save(tmp_path / "a.docx")
    docx_uri = store.create(SESSION, "media", "a.docx", (tmp_path / "a.docx").read_bytes()).uri()
    inspected = service.inspect(docx_uri)["inspection"]
    assert (inspected["paragraphs"], inspected["tables"]) == (2, 1)
    heading = inspected["headings"][0]
    blocks = service.read(docx_uri, str(heading["block"]))["blocks"]
    assert blocks[0]["text"] == "Informe anual"

    workbook = Workbook()
    workbook.active.title = "Datos"
    workbook.active["A1"] = 10
    workbook.active["A2"] = "=A1*2"
    workbook.save(tmp_path / "a.xlsx")
    xlsx_uri = store.create(SESSION, "media", "a.xlsx", (tmp_path / "a.xlsx").read_bytes()).uri()
    sheet = service.inspect(xlsx_uri)["inspection"]["sheets"][0]
    assert (sheet["name"], sheet["values"], sheet["formulas"]) == ("Datos", 2, 1)
    rows = service.read(xlsx_uri, "Datos!A1:A2")["rows"]
    assert rows[1]["cells"][0] == {"cell": "A2", "value": "=A1*2", "formula": True, "cached": None}
    with pytest.raises(DocumentError) as err:
        service.read(xlsx_uri, "Datos!A1:ZZ9999")
    assert err.value.code is DocumentErrorCode.DOCUMENT_LIMIT_EXCEEDED
