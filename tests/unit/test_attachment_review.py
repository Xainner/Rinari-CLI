"""Independent attachment acceptance checks using actual document parsers."""

import pytest
from PIL import Image

from rinari.artifacts.attachments import prepare_attachments
from rinari.artifacts.store import ArtifactStore
from rinari.models.images import references


def test_import_preserves_name_and_bytes_after_original_deleted(app_ctx, tmp_path):
    path = tmp_path / "captura con ñ.png"
    Image.new("RGB", (48, 48), "blue").save(path)
    original = path.read_bytes()
    store = ArtifactStore(app_ctx)
    prepared = prepare_attachments(store, "ses_review", [str(path)])
    path.unlink()
    item = prepared[0]
    assert item.name == "captura con ñ.png"
    assert store.get(item.source.uri()) == original
    refs = references(
        store, "ses_review", [{"uri": item.source.uri(), "sha256": item.source.sha256}]
    )
    assert refs[0].encoded()


def test_docx_paragraph_and_table_order(app_ctx, tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("Antes: mañana.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Nombre"
    table.cell(0, 1).text = "José"
    document.add_paragraph("Después: café.")
    path = tmp_path / "documento con ñ.docx"
    document.save(path)
    item = prepare_attachments(ArtifactStore(app_ctx), "ses_review", [str(path)])[0]
    text = item.extracted_text
    assert text.index("Antes") < text.index("José") < text.index("Después")


def test_xlsx_coordinates_formulas_and_multiple_sheets(app_ctx, tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Año"
    workbook.active["B2"] = "=SUM(1,2)"
    workbook.create_sheet("Resumen")["C3"] = "Información"
    path = tmp_path / "hojas con ñ.xlsx"
    workbook.save(path)
    item = prepare_attachments(ArtifactStore(app_ctx), "ses_review", [str(path)])[0]
    text = item.extracted_text
    assert all(
        value in text for value in ("Año", "B2", "=SUM(1,2)", "Resumen", "C3", "Información")
    )


def test_truncated_context_retains_full_derived_artifact(app_ctx, tmp_path):
    store = ArtifactStore(app_ctx)
    content = "información " * 12_000
    path = tmp_path / "extenso.txt"
    path.write_text(content, encoding="utf-8")
    item = prepare_attachments(store, "ses_review", [str(path)])[0]
    assert len(item.extracted_text) <= 64_000
    assert item.truncated
    assert item.derived is not None
    assert store.get(item.derived.uri()).decode("utf-8") == content


def test_corrupt_image_is_rejected_during_preparation(app_ctx, tmp_path):
    path = tmp_path / "corrupt.png"
    path.write_bytes(b"\x89PNG\r\n\x1a\nnot-an-image")
    with pytest.raises((ValueError, OSError)):
        prepare_attachments(ArtifactStore(app_ctx), "ses_review", [str(path)])


def test_reference_cannot_cross_sessions(app_ctx, tmp_path):
    store = ArtifactStore(app_ctx)
    path = tmp_path / "nota.txt"
    path.write_text("Datos adjuntos", encoding="utf-8")
    item = prepare_attachments(store, "ses_origin", [str(path)])[0]
    with pytest.raises(ValueError, match="session"):
        prepare_attachments(store, "ses_other", [item.reference()])


def test_cached_extraction_survives_deleted_source_without_reparsing(
    app_ctx, tmp_path, monkeypatch
):
    from rinari.artifacts import attachments

    store = ArtifactStore(app_ctx)
    path = tmp_path / "nota.txt"
    path.write_text("Mañana", encoding="utf-8")
    first = prepare_attachments(store, "ses_review", [str(path)])[0]
    path.unlink()

    def unexpected_read(_path):
        pytest.fail("Cached document was parsed again")

    monkeypatch.setattr(attachments, "_read_text", unexpected_read)
    second = prepare_attachments(store, "ses_review", [first.reference()])[0]
    assert second.extracted_text == "Mañana"
    assert second.name == "nota.txt"


def test_cancelled_preparation_does_not_import_files(app_ctx, tmp_path):
    from rinari.runtime.cancellation import CancellationToken
    from rinari.shared.errors import CancelledError

    token = CancellationToken()
    token.cancel()
    path = tmp_path / "nota.txt"
    path.write_text("Texto", encoding="utf-8")
    store = ArtifactStore(app_ctx)
    with pytest.raises(CancelledError):
        prepare_attachments(store, "ses_review", [str(path)], cancellation=token)
    assert not store.list(session_id="ses_review")


def test_pdf_partial_extraction_warning_survives_cache(app_ctx, tmp_path, monkeypatch):
    import pypdfium2 as pdfium

    from rinari.artifacts import attachments

    path = tmp_path / "more pages.pdf"
    document = pdfium.PdfDocument.new()
    for _ in range(21):
        document.new_page(80, 80).close()
    document.save(path)
    document.close()
    monkeypatch.setattr(attachments, "run_ocr", lambda *args, **kwargs: "Synthetic OCR")
    store = ArtifactStore(app_ctx)
    first = prepare_attachments(store, "ses_review", [str(path)])[0]
    assert first.warning and "20" in first.warning
    assert "[Page 21" not in first.extracted_text
    second = prepare_attachments(store, "ses_review", [first.reference()])[0]
    assert second.warning == first.warning


def test_context_links_to_derived_document_for_full_read(app_ctx, tmp_path):
    from rinari.artifacts.attachments import attachment_prompt

    path = tmp_path / "long.txt"
    path.write_text("Información\n" * 10_000, encoding="utf-8")
    prepared = prepare_attachments(ArtifactStore(app_ctx), "ses_review", [str(path)])
    assert prepared[0].derived.uri() in attachment_prompt(prepared)


def test_html_source_attachment_is_text_not_executed(app_ctx, tmp_path):
    path = tmp_path / "snake.html"
    content = '<html><script>throw new Error("attachment data only")</script></html>'
    path.write_text(content, encoding="utf-8")
    item = prepare_attachments(ArtifactStore(app_ctx), "ses_review", [str(path)])[0]
    assert item.kind == "text"
    assert item.extracted_text == content


def test_preparation_job_reconstructs_completed_attachment(app_ctx, tmp_path):
    import time

    from rinari.artifacts.attachments import AttachmentPreparationJobs

    path = tmp_path / "nota.txt"
    path.write_text("Información", encoding="utf-8")
    store = ArtifactStore(app_ctx)
    jobs = AttachmentPreparationJobs(store)
    job = jobs.start("ses_review", [str(path)])
    deadline = time.monotonic() + 5
    while job["status"] == "preparing" and time.monotonic() < deadline:
        time.sleep(0.01)
        job = jobs.get(job["job_id"])
    assert job["status"] == "ready", job
    path.unlink()
    recovered = AttachmentPreparationJobs(store).get(job["job_id"])
    assert recovered["status"] == "ready"
    assert recovered["attachments"] == job["attachments"]
    assert recovered["session_id"] == "ses_review"


def test_pdf_selected_visual_pages_are_real_persisted_images(app_ctx, tmp_path, monkeypatch):
    import pypdfium2 as pdfium

    from rinari.artifacts import attachments

    path = tmp_path / "selección visual.pdf"
    document = pdfium.PdfDocument.new()
    for _ in range(3):
        document.new_page(80, 80).close()
    document.save(path)
    document.close()
    monkeypatch.setattr(attachments, "run_ocr", lambda *a, **kw: "Texto sintético")
    store = ArtifactStore(app_ctx)
    item = prepare_attachments(
        store, "ses_review", [{"path": str(path), "page_range": "2-3", "visual_pages": [2]}]
    )[0]
    assert "[Page 1" not in item.extracted_text
    assert "[Page 2, OCR]" in item.extracted_text
    assert len(item.images) == 1
    image_ref = references(store, "ses_review", list(item.images))[0]
    assert image_ref.encoded()
    path.unlink()
    again = prepare_attachments(store, "ses_review", [item.reference()])[0]
    assert again.images == item.images
    assert again.warning == item.warning


def test_pdf_visual_page_outside_selected_range_is_rejected(app_ctx, tmp_path, monkeypatch):
    import pypdfium2 as pdfium

    from rinari.artifacts import attachments

    path = tmp_path / "rango.pdf"
    document = pdfium.PdfDocument.new()
    for _ in range(2):
        document.new_page(40, 40).close()
    document.save(path)
    document.close()
    monkeypatch.setattr(attachments, "run_ocr", lambda *a, **kw: "Texto")
    with pytest.raises(ValueError, match="range"):
        prepare_attachments(
            ArtifactStore(app_ctx),
            "ses_review",
            [{"path": str(path), "page_range": "1", "visual_pages": [2]}],
        )


@pytest.mark.parametrize(
    "vision,confirmed,accepted",
    [
        (True, False, True),
        (False, False, False),
        (False, True, False),
        (None, False, False),
        (None, True, True),
    ],
)
def test_engine_revalidates_vision_before_send(
    app_ctx, tmp_path, monkeypatch, vision, confirmed, accepted
):
    from types import SimpleNamespace

    from rinari.cli import agent_runtime
    from rinari.engine_protocol.errors import EngineProtocolError
    from rinari.engine_protocol.server import _prepare_turn_attachments

    path = tmp_path / "vision.png"
    Image.new("RGB", (24, 24), "red").save(path)
    store = ArtifactStore(app_ctx)
    prepared = prepare_attachments(store, "ses_review", [str(path)])[0]
    services = SimpleNamespace(artifacts=store, sessions=SimpleNamespace(show=lambda _: object()))
    monkeypatch.setattr(
        agent_runtime,
        "_caller_for",
        lambda *args: SimpleNamespace(capabilities=lambda: SimpleNamespace(vision=vision)),
    )
    if accepted:
        _, images, metadata = _prepare_turn_attachments(
            services, "ses_review", [prepared.reference()], allow_unconfirmed_vision=confirmed
        )
        assert images[0]["uri"] == prepared.source.uri()
        assert metadata[0]["name"] == "vision.png"
    else:
        with pytest.raises(EngineProtocolError, match="Vision"):
            _prepare_turn_attachments(
                services, "ses_review", [prepared.reference()], allow_unconfirmed_vision=confirmed
            )


def test_ocr_interrupt_reaps_its_process(monkeypatch, tmp_path):
    from rinari.artifacts import ocr

    class Process:
        returncode = None
        killed = False
        reaped = False

        def poll(self):
            return self.returncode

        def kill(self):
            self.killed = True
            self.returncode = -9

        def communicate(self, **kwargs):
            if not self.killed:
                raise KeyboardInterrupt
            self.reaped = True
            return "", ""

    process = Process()
    monkeypatch.setattr(ocr, "resolve_tesseract", lambda: (tmp_path / "tesseract.exe", None))
    monkeypatch.setattr(ocr.subprocess, "Popen", lambda *a, **kw: process)
    with pytest.raises(KeyboardInterrupt):
        ocr.run_ocr(tmp_path / "image.png")
    assert process.killed and process.reaped
