"""Compilador de ReportSpec a DOCX editable con python-docx.

Estilos de verdad (Title, Heading 1-3, List Bullet/Number, Caption), tablas
con cabecera repetida en cada página, pies de tabla y figura numerados como
campos SEQ, número de página como campo PAGE y un índice como campo TOC. Los
campos que dependen de la paginación (el índice) solo quedan actualizados si
Word los recalcula: eso se informa, no se supone.
"""

from __future__ import annotations

import contextlib
import datetime as dt
import io
from dataclasses import dataclass, field
from typing import Any

from rinari.documents.adapters import report_spec
from rinari.documents.contracts import DocumentError, DocumentErrorCode
from rinari.documents.design.tokens import Theme, theme

_TONE_FILL = {"info": "EEF3FE", "warning": "FFF6D6", "success": "E8F6EF"}
_TONE_BAR = {"info": None, "warning": "C98A00", "success": None}


@dataclass(slots=True)
class BuildPlan:
    headings: int = 0
    tables: int = 0
    figures: int = 0
    toc: bool = False
    fields: list[str] = field(default_factory=list)
    findings: list[dict[str, Any]] = field(default_factory=list)


def _rgb(value: str):
    from docx.shared import RGBColor

    return RGBColor.from_string(value.lstrip("#").upper())


def _set_cell_fill(cell, color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    props = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:val"), "clear")
    shade.set(qn("w:color"), "auto")
    shade.set(qn("w:fill"), color)
    props.append(shade)


def _table_borders(table, color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)
        borders.append(node)
    props.append(borders)


def _repeat_header(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    props = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    props.append(header)
    no_split = OxmlElement("w:cantSplit")
    props.append(no_split)


def _keep_row(row) -> None:
    """Una fila (un aviso, una tarjeta) no se parte entre páginas."""
    from docx.oxml import OxmlElement

    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _field(paragraph, instruction: str, result: str) -> None:
    """Un campo complejo con su resultado ya escrito (fldChar begin/sep/end)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def char(kind: str):
        run = paragraph.add_run()
        node = OxmlElement("w:fldChar")
        node.set(qn("w:fldCharType"), kind)
        run._r.append(node)
        return run

    char("begin")
    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run._r.append(instr)
    char("separate")
    paragraph.add_run(result)
    char("end")


class _Builder:
    def __init__(self, spec: dict[str, Any], resources: dict[str, str]) -> None:
        from docx import Document

        self.spec = report_spec.validate(spec)
        self.theme: Theme = theme(spec.get("theme"))
        self.labels = report_spec.labels(spec.get("language"))
        self.language = str(spec.get("language") or "es")
        self.resources = resources
        self.doc = Document()
        self.plan = BuildPlan()
        self.tables = 0
        self.figures = 0
        self.appendices = 0

    # -- estilos ---------------------------------------------------------------------
    def _styles(self) -> None:
        from docx.shared import Pt

        t = self.theme
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = t.body_font
        normal.font.size = Pt(11)
        normal.font.color.rgb = _rgb(t.text if not t.dark else "1F2937")
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        rpr = normal.element.get_or_add_rPr()
        fonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
        if fonts is not None:
            for attribute in ("w:eastAsia", "w:cs"):
                from docx.oxml.ns import qn

                fonts.set(qn(attribute), t.body_font)
        sizes = {1: 18, 2: 14, 3: 12}
        accent = t.accent if not t.dark else "2F5BEA"
        for level, size in sizes.items():
            style = styles[f"Heading {level}"]
            style.font.name = t.heading_font
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = _rgb(accent if level == 1 else "1F2937")
            style.paragraph_format.space_before = Pt(18 if level == 1 else 12)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = True
        title = styles["Title"]
        title.font.name = t.heading_font
        title.font.size = Pt(30)
        title.font.color.rgb = _rgb("111827")
        caption = styles["Caption"]
        caption.font.size = Pt(9)
        caption.font.italic = False
        caption.font.color.rgb = _rgb("5A6372")

    def _page(self) -> None:
        from docx.enum.section import WD_ORIENT
        from docx.shared import Cm

        page = self.spec.get("page") or {}
        width, height = report_spec.PAGE_SIZES[page.get("size", "A4")]
        section = self.doc.sections[0]
        if page.get("orientation") == "landscape":
            section.orientation = WD_ORIENT.LANDSCAPE
            width, height = height, width
        section.page_width = Cm(width)
        section.page_height = Cm(height)
        margin = Cm(float(page.get("margins_cm", 2.5)))
        section.left_margin = section.right_margin = margin
        section.top_margin = section.bottom_margin = margin
        self.text_width_cm = width - 2 * float(page.get("margins_cm", 2.5))
        cover = self.spec.get("cover", bool(self.spec.get("title")))
        section.different_first_page_header_footer = bool(cover)
        if self.spec.get("header"):
            paragraph = section.header.paragraphs[0]
            paragraph.text = str(self.spec["header"])
            paragraph.style = self.doc.styles["Header"]
        footer = section.footer.paragraphs[0]
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        if self.spec.get("footer"):
            footer.add_run(str(self.spec["footer"]) + "    ")
        footer.add_run(f"{self.labels['page']} ")
        _field(footer, "PAGE", "1")
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self.plan.fields.append("PAGE")

    # -- bloques -----------------------------------------------------------------------
    def _text(
        self, paragraph, text: str, *, bold: bool = False, color: str | None = None, size=None
    ):
        for piece, strong, italic in report_spec.runs(text):
            run = paragraph.add_run(piece)
            run.bold = bold or strong or None
            run.italic = italic or None
            if color:
                run.font.color.rgb = _rgb(color)
            if size:
                run.font.size = size
        return paragraph

    def _cover(self) -> None:
        from docx.shared import Pt

        spec = self.spec
        for _ in range(6):
            self.doc.add_paragraph()
        title = self.doc.add_paragraph(style="Title")
        title.add_run(str(spec["title"]))
        if spec.get("subtitle"):
            sub = self.doc.add_paragraph()
            self._text(sub, str(spec["subtitle"]), color="5A6372", size=Pt(14))
        meta = " · ".join(str(v) for v in (spec.get("author"), spec.get("date")) if v)
        if meta:
            line = self.doc.add_paragraph()
            line.paragraph_format.space_before = Pt(24)
            self._text(line, meta, color="5A6372")
        self.doc.add_page_break()

    def _toc(self) -> None:
        # Título con aspecto de capítulo pero sin nivel de esquema: no entra en el índice.
        from docx.shared import Pt

        heading = self.doc.add_paragraph()
        self._text(
            heading,
            self.labels["contents"],
            bold=True,
            color=self.theme.accent if not self.theme.dark else "2F5BEA",
            size=Pt(18),
        )
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_after = Pt(10)
        paragraph = self.doc.add_paragraph()
        _field(paragraph, 'TOC \\o "1-3" \\h \\z \\u', self.labels["toc_hint"])
        self.plan.toc = True
        self.plan.fields.append("TOC")
        self.doc.add_page_break()

    def _caption(self, kind: str, text: str | None) -> None:
        label = self.labels[kind]
        number = self.tables if kind == "table" else self.figures
        paragraph = self.doc.add_paragraph(style="Caption")
        paragraph.add_run(f"{label} ")
        _field(paragraph, f"SEQ {label} \\* ARABIC", str(number))
        if text:
            paragraph.add_run(f". {report_spec.plain(text)}")

    def _table(self, block: dict[str, Any]) -> None:
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Pt

        from rinari.documents.adapters.pptx_build import _decimals, _format, _is_number

        t = self.theme
        columns = [str(c) for c in block["columns"]]
        rows = block["rows"]
        self.tables += 1
        self.plan.tables += 1
        if block.get("caption"):
            self._caption("table", block["caption"])
        table = self.doc.add_table(rows=1, cols=len(columns))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(table, "D9DDE3")
        numeric = [
            bool(rows)
            and all(_is_number(r[c]) for r in rows if c < len(r) and r[c] not in (None, ""))
            for c in range(len(columns))
        ]
        decimals = [
            max(
                (_decimals(r[c]) for r in rows if c < len(r) and isinstance(r[c], float)), default=0
            )
            for c in range(len(columns))
        ]
        header = table.rows[0]
        _repeat_header(header)
        accent = t.accent if not t.dark else "2F5BEA"
        for col, text in enumerate(columns):
            cell = header.cells[col]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(text)
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = _rgb("FFFFFF")
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.RIGHT if numeric[col] else WD_ALIGN_PARAGRAPH.LEFT
            )
            _set_cell_fill(cell, accent)
        highlight = block.get("highlight_row")
        for index, values in enumerate(rows):
            cells = table.add_row().cells
            for col in range(len(columns)):
                value = values[col] if col < len(values) else None
                paragraph = cells[col].paragraphs[0]
                run = paragraph.add_run(
                    _format(value, self.language, decimals[col]) if value is not None else ""
                )
                run.font.size = Pt(10)
                if highlight == index:
                    run.bold = True
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.RIGHT if numeric[col] else WD_ALIGN_PARAGRAPH.LEFT
                )
                paragraph.paragraph_format.space_after = Pt(0)
                if index % 2 == 1:
                    _set_cell_fill(cells[col], "F3F4F6")
        widths = block.get("widths")
        if isinstance(widths, list) and len(widths) == len(columns):
            total = sum(float(w) for w in widths) or 1
            for row in table.rows:
                for col, cell in enumerate(row.cells):
                    cell.width = Cm(self.text_width_cm * float(widths[col]) / total)
        if block.get("source"):
            note = self.doc.add_paragraph(style="Caption")
            note.add_run(f"{self.labels['source']}: {report_spec.plain(block['source'])}")
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def _image(self, block: dict[str, Any]) -> None:
        from docx.shared import Cm

        path = self.resources.get(str(block["image"]))
        if path is None:
            raise DocumentError(
                DocumentErrorCode.NOT_FOUND, f"Image {block['image']!r} not provided"
            )
        width = min(float(block.get("width_cm") or self.text_width_cm), self.text_width_cm)
        self.doc.add_picture(path, width=Cm(width))
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.figures += 1
        self.plan.figures += 1
        self._caption("figure", block.get("caption"))

    def _callout(self, block: dict[str, Any]) -> None:
        from docx.shared import Pt

        table = self.doc.add_table(rows=1, cols=1)
        _keep_row(table.rows[0])
        cell = table.rows[0].cells[0]
        tone = block.get("tone", "info")
        _set_cell_fill(cell, _TONE_FILL[tone])
        paragraph = cell.paragraphs[0]
        if block.get("title"):
            self._text(paragraph, str(block["title"]), bold=True)
            paragraph.paragraph_format.keep_with_next = True
            paragraph = cell.add_paragraph()
        self._text(paragraph, str(block["text"]))
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def _kpis(self, block: dict[str, Any]) -> None:
        from docx.shared import Pt

        t = self.theme
        items = block["items"]
        table = self.doc.add_table(rows=1, cols=len(items))
        _keep_row(table.rows[0])
        for col, item in enumerate(items):
            cell = table.rows[0].cells[col]
            _set_cell_fill(cell, "F3F4F6")
            label = cell.paragraphs[0]
            self._text(label, str(item.get("label") or ""), color="5A6372", size=Pt(9))
            value = cell.add_paragraph()
            self._text(value, str(item.get("value") or ""), bold=True, size=Pt(18))
            if item.get("delta"):
                delta = cell.add_paragraph()
                text = str(item["delta"])
                negative = text.strip().startswith(("-", "\u2212"))
                color = t.negative if negative else t.positive
                if t.dark:
                    color = "C8363D" if negative else "0E7A5F"
                self._text(delta, text, color=color, size=Pt(9), bold=True)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def build(self) -> bytes:
        from docx.shared import Pt

        self._styles()
        self._page()
        spec = self.spec
        if spec.get("cover", bool(spec.get("title"))) and spec.get("title"):
            self._cover()
        if spec.get("toc"):
            self._toc()
        for block in spec["blocks"]:
            kind = block["type"]
            if kind == "heading":
                self.doc.add_paragraph(
                    report_spec.plain(block["text"]), style=f"Heading {block.get('level', 1)}"
                )
                self.plan.headings += 1
            elif kind == "paragraph":
                self._text(self.doc.add_paragraph(), str(block["text"]))
            elif kind == "bullets":
                style = "List Number" if block.get("numbered") else "List Bullet"
                for item in block["items"]:
                    self._text(self.doc.add_paragraph(style=style), str(item))
            elif kind == "table":
                self._table(block)
            elif kind == "image":
                self._image(block)
            elif kind == "quote":
                paragraph = self.doc.add_paragraph(style="Quote")
                self._text(paragraph, str(block["text"]))
                if block.get("author"):
                    author = self.doc.add_paragraph()
                    self._text(author, f"— {block['author']}", color="5A6372", size=Pt(10))
            elif kind == "callout":
                self._callout(block)
            elif kind == "kpis":
                self._kpis(block)
            elif kind == "page_break":
                self.doc.add_page_break()
            elif kind == "appendix":
                self.doc.add_page_break()
                letter = report_spec.appendix_letter(self.appendices)
                self.appendices += 1
                self.doc.add_paragraph(
                    f"{self.labels['appendix']} {letter}. {report_spec.plain(block['title'])}",
                    style="Heading 1",
                )
                self.plan.headings += 1
        props = self.doc.core_properties
        props.title = str(spec.get("title") or "")[:250]
        props.author = str(spec.get("author") or "")[:120]
        props.language = self.language
        props.comments = "Generated by Rinari"
        if spec.get("date"):
            with contextlib.suppress(ValueError):
                props.created = dt.datetime.fromisoformat(str(spec["date"]))
        buffer = io.BytesIO()
        self.doc.save(buffer)
        return buffer.getvalue()


def build(spec: dict[str, Any], resources: dict[str, str] | None = None) -> tuple[bytes, BuildPlan]:
    builder = _Builder(spec, resources or {})
    data = builder.build()
    return data, builder.plan


def template_catalog() -> dict[str, Any]:
    return {
        "report": (
            "title, subtitle?, author?, date?, language, theme, page {size A4|Letter, "
            "orientation, margins_cm}, cover?, toc?, header?, footer?, blocks[]"
        ),
        "blocks": {
            "heading": "text, level 1-3",
            "paragraph": "text (inline **bold** and *italic*)",
            "bullets": "items[], numbered?",
            "table": (
                "columns[], rows[[...]], caption?, widths? (relative), highlight_row?, source?"
            ),
            "image": "image (resource name), width_cm?, caption?",
            "quote": "text, author?",
            "callout": "text, title?, tone info|warning|success",
            "kpis": "items[1-4] {label, value, delta?}",
            "page_break": "",
            "appendix": "title (starts a new page, lettered A, B, …)",
        },
        "outputs": (
            "kind docx (editable Word) or pdf (ReportLab, selectable text) from the same spec"
        ),
    }


__all__ = ["BuildPlan", "build", "template_catalog"]
