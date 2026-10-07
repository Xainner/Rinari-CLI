"""Edición tipada de libros existentes.

Dos rutas, elegidas por la operación y no por comodidad:

- **Parche OOXML de celdas** (`xlsx.set_cells`): toca solo la hoja afectada,
  el `calcPr` del libro y, si cambian fórmulas, retira `calcChain.xml`. El
  resto del paquete se copia byte a byte, así que `preserve_strict` se
  cumple de verdad en libros con gráficos, pivots o macros.
- **openpyxl** para cambios de estructura y formato (hojas nuevas, formatos,
  anchos, validaciones, formato condicional, paneles). openpyxl no conserva
  todo lo que lee (gráficos, imágenes, slicers…): estas operaciones solo
  corren con `preserve_best_effort` o `rebuild`, y el diff de preservación
  dice qué se perdió.

Una fórmula escrita queda sin resultado en caché hasta que un backend
certificado (Excel) calcule el libro: no se inventa el valor.
"""

from __future__ import annotations

import posixpath
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from rinari.documents.contracts import DocumentError, DocumentErrorCode

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}
_M = "{{{}}}".format(NS["m"])
_CELL = re.compile(r"^\$?([A-Z]{1,3})\$?(\d{1,7})$")
_RANGE = re.compile(r"^\$?[A-Z]{1,3}\$?\d{1,7}(:\$?[A-Z]{1,3}\$?\d{1,7})?$")
MAX_CELLS = 20_000

CELL_OPS = ("xlsx.set_cells",)
STRUCTURE_OPS = (
    "xlsx.add_sheet",
    "xlsx.set_format",
    "xlsx.set_column_width",
    "xlsx.add_validation",
    "xlsx.add_conditional_format",
    "xlsx.freeze",
)
OPERATIONS = CELL_OPS + STRUCTURE_OPS
_FIELDS: dict[str, set[str]] = {
    "xlsx.set_cells": {"sheet", "cells"},
    "xlsx.add_sheet": {"name", "rows", "after"},
    "xlsx.set_format": {"sheet", "range", "number_format", "bold", "fill", "font_color"},
    "xlsx.set_column_width": {"sheet", "widths"},
    "xlsx.add_validation": {"sheet", "range", "list", "between", "type"},
    "xlsx.add_conditional_format": {"sheet", "range", "rule", "count"},
    "xlsx.freeze": {"sheet", "cell"},
}


@dataclass(slots=True)
class EditState:
    allowed: set[str] = field(default_factory=set)
    allow_new: set[str] = field(default_factory=set)
    changes: list[dict[str, Any]] = field(default_factory=list)
    findings: list[dict[str, Any]] = field(default_factory=list)
    formulas_changed: bool = False
    structural: bool = False


def _invalid(message: str, **details: Any) -> DocumentError:
    return DocumentError(DocumentErrorCode.INVALID_SPEC, message, details=details)


def _conflict(message: str, **details: Any) -> DocumentError:
    return DocumentError(
        DocumentErrorCode.REVISION_CONFLICT,
        message,
        details=details,
        action="Read the range again and retry with current values",
    )


def validate(operations: Any) -> list[dict[str, Any]]:
    if not isinstance(operations, list) or not operations:
        raise _invalid("operations must be a non-empty list")
    total = 0
    for index, op in enumerate(operations, start=1):
        if not isinstance(op, dict) or op.get("op") not in _FIELDS:
            raise _invalid(f"operation {index}: unknown op", ops=list(OPERATIONS))
        unknown = sorted(set(op) - {"op"} - _FIELDS[op["op"]])
        if unknown:
            raise _invalid(
                f"operation {index} ({op['op']}): unknown fields {', '.join(unknown)}",
                allowed=sorted(_FIELDS[op["op"]]),
            )
        if op["op"] == "xlsx.set_cells":
            cells = op.get("cells")
            if not isinstance(cells, list) or not cells:
                raise _invalid(f"operation {index}: cells is a list of {{cell, value|formula}}")
            total += len(cells)
            for cell in cells:
                if not isinstance(cell, dict) or not _CELL.match(str(cell.get("cell") or "")):
                    raise _invalid(f"operation {index}: bad cell {cell!r}")
                extra = set(cell) - {"cell", "value", "formula", "expected"}
                if extra:
                    raise _invalid(f"operation {index}: unknown cell fields {sorted(extra)}")
                if "value" in cell and "formula" in cell:
                    raise _invalid(f"{cell['cell']}: value or formula, not both")
    if total > MAX_CELLS:
        raise DocumentError(
            DocumentErrorCode.DOCUMENT_LIMIT_EXCEEDED, f"At most {MAX_CELLS} cells per edit"
        )
    return operations


def needs_rebuild(operations: list[dict[str, Any]]) -> bool:
    return any(op["op"] in STRUCTURE_OPS for op in operations)


# -- paquete ------------------------------------------------------------------------------
class _Package:
    def __init__(self, path: Path) -> None:
        with zipfile.ZipFile(path) as archive:
            self.infos = archive.infolist()
            self.parts = {info.filename: archive.read(info.filename) for info in self.infos}
        self.changed: set[str] = set()
        self.removed: set[str] = set()

    def xml(self, name: str):
        from lxml import etree

        parser = etree.XMLParser(resolve_entities=False, no_network=True, remove_blank_text=False)
        return etree.fromstring(self.parts[name], parser)

    def put(self, name: str, root) -> None:
        from lxml import etree

        self.parts[name] = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        self.changed.add(name)

    def drop(self, name: str) -> None:
        if name in self.parts:
            del self.parts[name]
            self.removed.add(name)

    def save(self, target: Path) -> None:
        with zipfile.ZipFile(target, "w") as out:
            for info in self.infos:
                if info.filename in self.removed:
                    continue
                if info.filename in self.changed:
                    out.writestr(info.filename, self.parts[info.filename], zipfile.ZIP_DEFLATED)
                else:
                    # Lo que no se toca viaja igual: misma compresión, mismos bytes.
                    out.writestr(info, self.parts[info.filename])

    def rels_of(self, part: str) -> tuple[str, Any]:
        directory, name = posixpath.split(part)
        rels = posixpath.join(directory, "_rels", f"{name}.rels")
        return rels, self.xml(rels) if rels in self.parts else None

    def target(self, part: str, rel_target: str) -> str:
        if rel_target.startswith("/"):
            return rel_target.lstrip("/")
        return posixpath.normpath(posixpath.join(posixpath.dirname(part), rel_target))


def _sheet_part(package: _Package, sheet: str) -> str:
    workbook = package.xml("xl/workbook.xml")
    _, rels = package.rels_of("xl/workbook.xml")
    for node in workbook.iter(f"{_M}sheet"):
        if node.get("name") == sheet:
            rid = node.get("{{{}}}id".format(NS["r"]))
            for rel in rels:
                if rel.get("Id") == rid:
                    return package.target("xl/workbook.xml", rel.get("Target"))
    raise DocumentError(DocumentErrorCode.NOT_FOUND, f"No sheet named {sheet!r}")


def _split(address: str) -> tuple[str, int]:
    match = _CELL.match(address)
    assert match
    return match.group(1), int(match.group(2))


def _col_index(letters: str) -> int:
    value = 0
    for char in letters:
        value = value * 26 + ord(char) - 64
    return value


def _in_range(address: str, ref: str) -> bool:
    col, row = _split(address)
    first, _, last = ref.partition(":")
    last = last or first
    c1, r1 = _split(first.replace("$", ""))
    c2, r2 = _split(last.replace("$", ""))
    return _col_index(c1) <= _col_index(col) <= _col_index(c2) and r1 <= row <= r2


def _protected_ranges(package: _Package, part: str) -> tuple[list[str], list[str]]:
    """Cabeceras de tablas y celdas combinadas no ancla de la hoja."""
    headers: list[str] = []
    _, rels = package.rels_of(part)
    if rels is not None:
        for rel in rels:
            if rel.get("Type", "").endswith("/table"):
                table_part = package.target(part, rel.get("Target"))
                if table_part in package.parts:
                    table = package.xml(table_part)
                    ref = table.get("ref", "")
                    header_rows = int(table.get("headerRowCount", "1"))
                    if header_rows and ref:
                        first, _, last = ref.partition(":")
                        c1, r1 = _split(first)
                        c2, _ = _split(last or first)
                        headers.append(f"{c1}{r1}:{c2}{r1}")
    merged: list[str] = []
    sheet = package.xml(part)
    for node in sheet.iter(f"{_M}mergeCell"):
        ref = node.get("ref", "")
        if ":" in ref:
            merged.append(ref)
    return headers, merged


def _cell_text(cell, shared: list[str]) -> Any:
    kind = cell.get("t")
    formula = cell.find(f"{_M}f")
    if formula is not None:
        return "=" + (formula.text or "")
    value = cell.find(f"{_M}v")
    if kind == "inlineStr":
        return "".join(t.text or "" for t in cell.iter(f"{_M}t"))
    if value is None or value.text is None:
        return None
    if kind == "s":
        index = int(value.text)
        return shared[index] if index < len(shared) else None
    if kind == "b":
        return value.text == "1"
    if kind in ("str", "e"):
        return value.text
    try:
        number = float(value.text)
        return int(number) if number.is_integer() else number
    except ValueError:
        return value.text


def _shared_strings(package: _Package) -> list[str]:
    name = "xl/sharedStrings.xml"
    if name not in package.parts:
        return []
    root = package.xml(name)
    return ["".join(t.text or "" for t in si.iter(f"{_M}t")) for si in root.iter(f"{_M}si")]


def _same(current: Any, expected: Any) -> bool:
    if isinstance(expected, (int, float)) and isinstance(current, (int, float)):
        return abs(float(current) - float(expected)) < 1e-9
    return (current if current is not None else None) == expected or str(current) == str(expected)


def _find_row(sheet_data, number: int):
    from lxml import etree

    rows = sheet_data.findall(f"{_M}row")
    for row in rows:
        r = int(row.get("r", "0"))
        if r == number:
            return row
        if r > number:
            new = etree.Element(f"{_M}row", r=str(number))
            row.addprevious(new)
            return new
    new = etree.SubElement(sheet_data, f"{_M}row", r=str(number))
    return new


def _find_cell(row, address: str):
    from lxml import etree

    letters, _ = _split(address)
    target = _col_index(letters)
    for cell in row.findall(f"{_M}c"):
        ref = cell.get("r")
        col = _col_index(_split(ref)[0]) if ref else 0
        if col == target:
            return cell
        if col > target:
            new = etree.Element(f"{_M}c", r=address)
            cell.addprevious(new)
            return new
    return etree.SubElement(row, f"{_M}c", r=address)


def _clear(cell) -> None:
    for child in list(cell):
        cell.remove(child)
    cell.attrib.pop("t", None)


def _set_cells(package: _Package, op: dict[str, Any], state: EditState) -> None:
    from lxml import etree

    sheet_name = str(op.get("sheet") or "")
    part = _sheet_part(package, sheet_name)
    headers, merged = _protected_ranges(package, part)
    shared = _shared_strings(package)
    root = package.xml(part)
    data = root.find(f"{_M}sheetData")
    if data is None:
        raise DocumentError(DocumentErrorCode.UNSUPPORTED_FEATURE, "Sheet without data")
    touched = []
    for item in op["cells"]:
        address = item["cell"].replace("$", "")
        if any(_in_range(address, ref) for ref in headers):
            raise DocumentError(
                DocumentErrorCode.UNSUPPORTED_FEATURE,
                f"{sheet_name}!{address} is a table header; renaming columns is not supported",
            )
        for ref in merged:
            if _in_range(address, ref) and address != ref.split(":")[0].replace("$", ""):
                raise DocumentError(
                    DocumentErrorCode.UNSUPPORTED_FEATURE,
                    f"{sheet_name}!{address} is inside a merged range; write its first cell",
                )
        _, number = _split(address)
        row = _find_row(data, number)
        row.attrib.pop("spans", None)
        cell = _find_cell(row, address)
        formula = cell.find(f"{_M}f")
        if formula is not None and formula.get("t") in ("array", "dataTable"):
            raise DocumentError(
                DocumentErrorCode.UNSUPPORTED_FEATURE, f"{address} is part of an array formula"
            )
        if formula is not None and formula.get("t") == "shared" and formula.get("ref"):
            raise DocumentError(
                DocumentErrorCode.UNSUPPORTED_FEATURE,
                f"{address} anchors a shared formula used by {formula.get('ref')}",
            )
        current = _cell_text(cell, shared)
        if "expected" in item and not _same(current, item["expected"]):
            raise _conflict(f"{sheet_name}!{address} changed", current=current)
        if formula is not None:
            state.formulas_changed = True
        _clear(cell)
        if "formula" in item:
            text = str(item["formula"]).strip()
            if not text:
                raise _invalid(f"{address}: empty formula")
            etree.SubElement(cell, f"{_M}f").text = text.lstrip("=")
            state.formulas_changed = True
        else:
            value = item.get("value")
            if value is None:
                pass
            elif isinstance(value, bool):
                cell.set("t", "b")
                etree.SubElement(cell, f"{_M}v").text = "1" if value else "0"
            elif isinstance(value, (int, float)):
                etree.SubElement(cell, f"{_M}v").text = (
                    repr(float(value)) if isinstance(value, float) else str(value)
                )
            else:
                # Texto en línea: nunca se interpreta como fórmula y no toca sharedStrings.
                cell.set("t", "inlineStr")
                holder = etree.SubElement(cell, f"{_M}is")
                node = etree.SubElement(holder, f"{_M}t")
                node.text = str(value)
                if node.text != node.text.strip():
                    node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        if len(cell) == 0 and not cell.get("s"):
            cell.getparent().remove(cell)
        touched.append(address)
    _extend_dimension(root, touched)
    package.put(part, root)
    state.allowed.add(part)
    state.changes.append(
        {"op": "xlsx.set_cells", "sheet": sheet_name, "cells": touched[:200], "count": len(touched)}
    )


def _extend_dimension(root, addresses: list[str]) -> None:
    node = root.find(f"{_M}dimension")
    if node is None or not addresses:
        return
    ref = node.get("ref", "A1")
    first, _, last = ref.partition(":")
    last = last or first
    try:
        c1, r1 = _split(first)
        c2, r2 = _split(last)
    except AssertionError:
        return
    cols = [_col_index(c1), _col_index(c2)]
    rows = [r1, r2]
    for address in addresses:
        letters, number = _split(address)
        cols.append(_col_index(letters))
        rows.append(number)
    from rinari.documents.adapters.xlsx_build import column_letter

    node.set(
        "ref",
        f"{column_letter(min(cols) - 1)}{min(rows)}:{column_letter(max(cols) - 1)}{max(rows)}",
    )


def _stale_calculation(package: _Package, state: EditState) -> None:
    """Las cachés dejan de ser fiables: Excel recalcula al abrir y el calcChain se retira."""
    workbook = package.xml("xl/workbook.xml")
    calc = workbook.find(f"{_M}calcPr")
    if calc is None:
        from lxml import etree

        calc = etree.SubElement(workbook, f"{_M}calcPr")
    calc.set("fullCalcOnLoad", "1")
    package.put("xl/workbook.xml", workbook)
    state.allowed.add("xl/workbook.xml")
    chain = "xl/calcChain.xml"
    if chain in package.parts:
        package.drop(chain)
        state.allowed.add(chain)
        rels_name, rels = package.rels_of("xl/workbook.xml")
        for rel in list(rels):
            if rel.get("Target", "").endswith("calcChain.xml"):
                rels.remove(rel)
        package.put(rels_name, rels)
        types = package.xml("[Content_Types].xml")
        for node in list(types):
            if node.get("PartName") == "/xl/calcChain.xml":
                types.remove(node)
        package.put("[Content_Types].xml", types)


# -- openpyxl ------------------------------------------------------------------------------
def _structure(
    source: Path, target: Path, operations: list[dict[str, Any]], state: EditState
) -> None:
    import openpyxl
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, Rule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = openpyxl.load_workbook(str(source), keep_links=True, rich_text=True)

    def sheet(name: Any):
        if name not in workbook.sheetnames:
            raise DocumentError(DocumentErrorCode.NOT_FOUND, f"No sheet named {name!r}")
        return workbook[name]

    def check_range(ref: Any) -> str:
        if not isinstance(ref, str) or not _RANGE.match(ref):
            raise _invalid(f"Bad range {ref!r}")
        return ref

    for op in operations:
        kind = op["op"]
        if kind == "xlsx.set_cells":
            ws = sheet(op.get("sheet"))
            for item in op["cells"]:
                cell = ws[item["cell"]]
                if "expected" in item and not _same(cell.value, item["expected"]):
                    raise _conflict(f"{ws.title}!{item['cell']} changed", current=cell.value)
                if "formula" in item:
                    cell.value = "=" + str(item["formula"]).lstrip("=")
                    state.formulas_changed = True
                else:
                    value = item.get("value")
                    cell.value = value
                    if isinstance(value, str) and value.startswith("="):
                        cell.data_type = "s"
        elif kind == "xlsx.add_sheet":
            name = str(op.get("name") or "")
            if not re.match(r"^[^\[\]:*?/\\]{1,31}$", name) or name in workbook.sheetnames:
                raise _invalid(f"Bad or existing sheet name {name!r}")
            ws = workbook.create_sheet(name)
            for values in op.get("rows") or []:
                if not isinstance(values, list):
                    raise _invalid("rows is a list of lists")
                ws.append([v for v in values])
                for cell in ws[ws.max_row]:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        cell.data_type = "s"
        elif kind == "xlsx.set_format":
            ws = sheet(op.get("sheet"))
            for row in ws[check_range(op.get("range"))]:
                for cell in row if isinstance(row, tuple) else (row,):
                    if op.get("number_format"):
                        cell.number_format = str(op["number_format"])
                    if "bold" in op or op.get("font_color"):
                        cell.font = Font(
                            bold=bool(op.get("bold", cell.font.bold)),
                            color=str(
                                op.get("font_color")
                                or (cell.font.color.rgb if cell.font.color else None)
                                or "FF000000"
                            ),
                            name=cell.font.name,
                            size=cell.font.size,
                        )
                    if op.get("fill"):
                        cell.fill = PatternFill("solid", fgColor=str(op["fill"]).lstrip("#"))
        elif kind == "xlsx.set_column_width":
            ws = sheet(op.get("sheet"))
            for column, width in (op.get("widths") or {}).items():
                if not re.match(r"^[A-Z]{1,3}$", str(column)):
                    raise _invalid(f"Bad column {column!r}")
                ws.column_dimensions[str(column)].width = float(width)
        elif kind == "xlsx.add_validation":
            ws = sheet(op.get("sheet"))
            ref = check_range(op.get("range"))
            if isinstance(op.get("list"), list):
                items = ",".join(str(v).replace('"', "'") for v in op["list"])
                validation = DataValidation(type="list", formula1=f'"{items}"', allow_blank=True)
            elif isinstance(op.get("between"), list) and len(op["between"]) == 2:
                low, high = op["between"]
                validation = DataValidation(
                    type="decimal" if op.get("type", "decimal") == "decimal" else "whole",
                    operator="between",
                    formula1=str(low),
                    formula2=str(high),
                )
            else:
                raise _invalid("validation needs list or between")
            ws.add_data_validation(validation)
            validation.add(ref)
        elif kind == "xlsx.add_conditional_format":
            ws = sheet(op.get("sheet"))
            ref = check_range(op.get("range"))
            rule = op.get("rule")
            if rule == "negative_red":
                ws.conditional_formatting.add(
                    ref, CellIsRule(operator="lessThan", formula=["0"], font=Font(color="FFC8363D"))
                )
            elif rule == "data_bar":
                ws.conditional_formatting.add(
                    ref, DataBarRule(start_type="min", end_type="max", color="FF2F5BEA")
                )
            elif rule == "color_scale":
                ws.conditional_formatting.add(
                    ref,
                    ColorScaleRule(
                        start_type="min",
                        start_color="FFF8696B",
                        mid_type="percentile",
                        mid_value=50,
                        mid_color="FFFFEB84",
                        end_type="max",
                        end_color="FF63BE7B",
                    ),
                )
            elif rule == "top":
                ws.conditional_formatting.add(
                    ref, Rule(type="top10", rank=int(op.get("count") or 10), dxf=None)
                )
            else:
                raise _invalid("rule must be negative_red|data_bar|color_scale|top")
        elif kind == "xlsx.freeze":
            ws = sheet(op.get("sheet"))
            cell = str(op.get("cell") or "")
            if not _CELL.match(cell):
                raise _invalid("freeze cell like A2")
            ws.freeze_panes = cell
        state.changes.append(
            {k: v for k, v in op.items() if k not in ("rows", "cells")}
            | (
                {"count": len(op.get("cells") or op.get("rows") or [])}
                if op.get("cells") or op.get("rows")
                else {}
            )
        )
    if state.formulas_changed:
        workbook.calculation.fullCalcOnLoad = True
    workbook.save(str(target))
    state.structural = True


def apply(
    source: Path, target: Path, operations: list[dict[str, Any]], *, rebuild_allowed: bool
) -> EditState:
    validate(operations)
    state = EditState()
    if needs_rebuild(operations):
        if not rebuild_allowed:
            raise DocumentError(
                DocumentErrorCode.PRESERVATION_RISK,
                "Structure and format changes rewrite the workbook with openpyxl, which can "
                "drop charts, images and other objects",
                action="Use preserve_best_effort (the diff reports any loss) or rebuild",
                details={
                    "operations": [op["op"] for op in operations if op["op"] in STRUCTURE_OPS]
                },
            )
        _structure(source, target, operations, state)
        return state
    package = _Package(source)
    for op in operations:
        _set_cells(package, op, state)
    if state.formulas_changed:
        _stale_calculation(package, state)
    package.save(target)
    return state


# -- lectura de fórmulas y cachés -------------------------------------------------------------
ERRORS = ("#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")


def formulas(path: Path, *, limit: int = 50_000, calculated: bool = False) -> dict[str, Any]:
    """Inventario de fórmulas con su caché: presente, ausente o con error.

    `calculated` dice que el archivo lo acaba de guardar el backend certificado
    tras recalcular: solo entonces las cachés son `fresh`.
    """
    import warnings

    import openpyxl

    warnings.filterwarnings("ignore", module="openpyxl")

    with_formulas = openpyxl.load_workbook(str(path), read_only=True, data_only=False)
    with_values = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    count = cached = missing = 0
    errors: list[dict[str, Any]] = []
    truncated = False
    try:
        for ws_f, ws_v in zip(with_formulas.worksheets, with_values.worksheets, strict=False):
            for row_f, row_v in zip(ws_f.iter_rows(), ws_v.iter_rows(), strict=False):
                for cell_f, cell_v in zip(row_f, row_v, strict=False):
                    value = cell_f.value
                    # El tipo de la celda, no un texto que empiece por "=".
                    if getattr(cell_f, "data_type", None) == "f" and isinstance(value, str):
                        count += 1
                        cache = cell_v.value
                        if cache is None:
                            missing += 1
                        else:
                            cached += 1
                            if isinstance(cache, str) and cache in ERRORS and len(errors) < 200:
                                errors.append(
                                    {
                                        "cell": f"{ws_f.title}!{cell_f.coordinate}",
                                        "error": cache,
                                        "formula": value[:200],
                                    }
                                )
                    if count >= limit:
                        truncated = True
                        break
                if truncated:
                    break
            if truncated:
                break
        full_calc = bool(getattr(with_formulas.calculation, "fullCalcOnLoad", False))
    finally:
        with_formulas.close()
        with_values.close()
    if count == 0:
        state = "none"
    elif calculated and not missing:
        state = "fresh"
    elif missing == count:
        state = "missing"
    elif full_calc or missing:
        state = "stale"
    else:
        state = "unknown"
    return {
        "formulas": count,
        "cached": cached,
        "missing": missing,
        "errors": errors,
        "cache_state": state,
        "full_calc_on_load": full_calc,
        "truncated": truncated,
    }


__all__ = [
    "ERRORS",
    "OPERATIONS",
    "STRUCTURE_OPS",
    "EditState",
    "apply",
    "formulas",
    "needs_rebuild",
    "validate",
]
