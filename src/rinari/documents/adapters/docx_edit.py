"""Edición tipada de documentos Word existentes con python-docx.

Los bloques se numeran como en `documents.read` (párrafos y tablas del cuerpo
en orden, desde 1). Cambiar texto conserva el formato de los runs: un
reemplazo que cruza runs deja el formato alrededor de la coincidencia, y
reescribir un párrafo conserva el estilo del párrafo y el formato de su
primer run. Comentarios, campos, hipervínculos y lo que no se toca quedan
como estaban; el diff de preservación lo comprueba.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any

from rinari.documents.contracts import DocumentError, DocumentErrorCode

OPERATIONS = (
    "docx.replace_text",
    "docx.set_paragraph",
    "docx.insert_paragraph",
    "docx.insert_table",
    "docx.delete_block",
    "docx.set_cell",
    "docx.add_comment",
)
_FIELDS: dict[str, set[str]] = {
    "docx.replace_text": {"find", "replace", "expected_count", "include_headers"},
    "docx.set_paragraph": {"block", "text", "expected_text"},
    "docx.insert_paragraph": {"after", "text", "style"},
    "docx.insert_table": {"after", "columns", "rows", "style"},
    "docx.delete_block": {"block", "expected_text"},
    "docx.set_cell": {"block", "row", "col", "text", "expected_text"},
    "docx.add_comment": {"block", "text", "author"},
}


@dataclass(slots=True)
class EditState:
    allowed: set[str] = field(default_factory=lambda: {"word/document.xml"})
    allow_new: set[str] = field(default_factory=set)
    changes: list[dict[str, Any]] = field(default_factory=list)


def _invalid(message: str, **details: Any) -> DocumentError:
    return DocumentError(DocumentErrorCode.INVALID_SPEC, message, details=details)


def _conflict(message: str, **details: Any) -> DocumentError:
    return DocumentError(
        DocumentErrorCode.REVISION_CONFLICT,
        message,
        details=details,
        action="Read the document again and retry with current values",
    )


def validate(operations: Any) -> list[dict[str, Any]]:
    if not isinstance(operations, list) or not operations:
        raise _invalid("operations must be a non-empty list")
    if len(operations) > 500:
        raise DocumentError(DocumentErrorCode.DOCUMENT_LIMIT_EXCEEDED, "At most 500 operations")
    for index, op in enumerate(operations, start=1):
        if not isinstance(op, dict) or op.get("op") not in _FIELDS:
            raise _invalid(f"operation {index}: unknown op", ops=list(OPERATIONS))
        unknown = sorted(set(op) - {"op"} - _FIELDS[op["op"]])
        if unknown:
            raise _invalid(
                f"operation {index} ({op['op']}): unknown fields {', '.join(unknown)}",
                allowed=sorted(_FIELDS[op["op"]]),
            )
    return operations


def _blocks(document) -> list[Any]:
    from rinari.documents.adapters.basic_read import _blocks as blocks

    return list(blocks(document))


def _block(document, number: Any):
    blocks = _blocks(document)
    if not isinstance(number, int) or not 1 <= number <= len(blocks):
        raise _invalid(f"block must be 1..{len(blocks)}", block=number)
    return blocks[number - 1]


def _paragraph(document, number: Any):
    from docx.text.paragraph import Paragraph

    block = _block(document, number)
    if not isinstance(block, Paragraph):
        raise _invalid(f"block {number} is a table")
    return block


def _set_paragraph(paragraph, text: str) -> None:
    """Texto nuevo con el formato del primer run; nada de runs viejos entre medias."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)


def _replace_in_paragraph(paragraph, find: str, replace: str) -> int:
    from rinari.documents.adapters.pptx_edit import _replace_in_paragraph as replace_runs

    return replace_runs(paragraph, find, replace)


def _paragraphs(document, include_headers: bool):
    for block in _blocks(document):
        if hasattr(block, "rows"):
            for row in block.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
        else:
            yield block
    if include_headers:
        for section in document.sections:
            for part in (section.header, section.footer):
                yield from part.paragraphs


def _new_paragraph_after(anchor_element, document, text: str, style: str | None):
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    element = OxmlElement("w:p")
    anchor_element.addnext(element)
    paragraph = Paragraph(element, document)
    if style:
        if style not in [s.name for s in document.styles]:
            raise _invalid(f"Style {style!r} does not exist in this document")
        paragraph.style = document.styles[style]
    paragraph.add_run(text)
    return paragraph


def _anchor(document, after: Any):
    """El elemento tras el que se inserta: 0 es el principio del cuerpo."""
    if after == 0:
        body = document.element.body
        first = body[0] if len(body) else None
        if first is None:
            raise _invalid("The document is empty")
        return first, True
    block = _block(document, after)
    return block._element, False


def apply_operations(document, operations: list[dict[str, Any]], state: EditState) -> None:
    from docx.table import Table

    for index, op in enumerate(operations, start=1):
        kind = op["op"]
        try:
            if kind == "docx.replace_text":
                find = op.get("find")
                if not isinstance(find, str) or not find:
                    raise _invalid("find must be a non-empty string")
                total = 0
                for paragraph in _paragraphs(document, bool(op.get("include_headers"))):
                    total += _replace_in_paragraph(paragraph, find, str(op.get("replace", "")))
                expected = op.get("expected_count")
                if expected is not None and total != int(expected):
                    raise _conflict(f"Found {total} occurrences, expected {expected}", found=total)
                if total == 0:
                    raise _conflict(f"{find!r} does not appear in the document")
                if op.get("include_headers"):
                    for section in document.sections:
                        for part in (section.header, section.footer):
                            if not part.is_linked_to_previous:
                                state.allowed.add(str(part.part.partname).lstrip("/"))
                state.changes.append({"op": kind, "count": total})
            elif kind == "docx.set_paragraph":
                paragraph = _paragraph(document, op.get("block"))
                if "expected_text" in op and paragraph.text != op["expected_text"]:
                    raise _conflict("The paragraph changed", current=paragraph.text[:300])
                _set_paragraph(paragraph, str(op.get("text", "")))
                state.changes.append({"op": kind, "block": op["block"]})
            elif kind == "docx.insert_paragraph":
                anchor, at_start = _anchor(document, op.get("after"))
                if at_start:
                    from docx.oxml import OxmlElement
                    from docx.text.paragraph import Paragraph

                    element = OxmlElement("w:p")
                    anchor.addprevious(element)
                    paragraph = Paragraph(element, document)
                    if op.get("style"):
                        paragraph.style = document.styles[op["style"]]
                    paragraph.add_run(str(op.get("text", "")))
                else:
                    _new_paragraph_after(anchor, document, str(op.get("text", "")), op.get("style"))
                state.changes.append({"op": kind, "after": op.get("after")})
            elif kind == "docx.insert_table":
                columns, rows = op.get("columns"), op.get("rows")
                if not isinstance(columns, list) or not columns or not isinstance(rows, list):
                    raise _invalid("insert_table needs columns and rows")
                anchor, at_start = _anchor(document, op.get("after"))
                table = document.add_table(rows=1, cols=len(columns))
                style = op.get("style") or (
                    "Table Grid" if "Table Grid" in [s.name for s in document.styles] else None
                )
                if style:
                    table.style = document.styles[style]
                for col, header in enumerate(columns):
                    table.rows[0].cells[col].text = str(header)
                for values in rows:
                    cells = table.add_row().cells
                    for col in range(len(columns)):
                        value = values[col] if col < len(values) else ""
                        cells[col].text = "" if value is None else str(value)
                element = table._tbl
                element.getparent().remove(element)
                if at_start:
                    anchor.addprevious(element)
                else:
                    anchor.addnext(element)
                state.changes.append({"op": kind, "after": op.get("after"), "rows": len(rows)})
            elif kind == "docx.delete_block":
                block = _block(document, op.get("block"))
                text = block.text if not isinstance(block, Table) else ""
                if "expected_text" in op and text != op["expected_text"]:
                    raise _conflict("The block changed", current=text[:300])
                element = block._element
                if element.getparent() is not None and len(_blocks(document)) < 2:
                    raise _invalid("A document keeps at least one block")
                element.getparent().remove(element)
                state.changes.append({"op": kind, "block": op["block"]})
            elif kind == "docx.set_cell":
                block = _block(document, op.get("block"))
                if not isinstance(block, Table):
                    raise _invalid(f"block {op.get('block')} is not a table")
                row, col = op.get("row"), op.get("col")
                if not isinstance(row, int) or not isinstance(col, int):
                    raise _invalid("row and col are 0-based integers")
                if not (0 <= row < len(block.rows) and 0 <= col < len(block.columns)):
                    raise _invalid("Cell out of range")
                cell = block.cell(row, col)
                if "expected_text" in op and cell.text != op["expected_text"]:
                    raise _conflict("The cell changed", current=cell.text[:300])
                paragraphs = cell.paragraphs
                _set_paragraph(paragraphs[0], str(op.get("text", "")))
                for extra in paragraphs[1:]:
                    extra._element.getparent().remove(extra._element)
                state.changes.append({"op": kind, "block": op["block"], "cell": [row, col]})
            elif kind == "docx.add_comment":
                paragraph = _paragraph(document, op.get("block"))
                if not paragraph.runs:
                    raise _invalid("Comments attach to text; this paragraph is empty")
                document.add_comment(
                    paragraph.runs,
                    text=str(op.get("text") or ""),
                    author=str(op.get("author") or "Rinari")[:60],
                    initials="R",
                )
                state.allowed.update({"word/comments.xml", "[Content_Types].xml"})
                state.allow_new.update({"word/comments", "word/_rels/"})
                state.changes.append({"op": kind, "block": op["block"]})
        except DocumentError as exc:
            exc.details = {**(exc.details or {}), "operation": index, "op": kind}
            raise


def apply(source, target, operations: list[dict[str, Any]]) -> EditState:
    from docx import Document

    validate(operations)
    document = Document(str(source))
    state = EditState()
    apply_operations(document, operations, state)
    document.save(str(target))
    return state


def render_template(
    template, target, context: dict[str, Any], *, autoescape: bool = True
) -> list[str]:
    """Una plantilla docxtpl con un entorno Jinja aislado y variables estrictas."""
    from docxtpl import DocxTemplate
    from jinja2 import StrictUndefined
    from jinja2.exceptions import SecurityError, TemplateError, UndefinedError
    from jinja2.sandbox import ImmutableSandboxedEnvironment

    if not isinstance(context, dict):
        raise _invalid("context must be an object")
    template_doc = DocxTemplate(str(template))
    environment = ImmutableSandboxedEnvironment(undefined=StrictUndefined, autoescape=False)
    try:
        variables = sorted(template_doc.get_undeclared_template_variables(environment))
        template_doc.render(context, jinja_env=environment, autoescape=autoescape)
    except UndefinedError as exc:
        raise DocumentError(
            DocumentErrorCode.INVALID_SPEC,
            f"The template needs a value that the context does not have: {exc}",
        ) from exc
    except SecurityError as exc:
        raise DocumentError(
            DocumentErrorCode.UNSAFE_EXTERNAL_RESOURCE,
            f"The template tried an unsafe operation: {exc}",
        ) from exc
    except TemplateError as exc:
        raise DocumentError(DocumentErrorCode.INVALID_SPEC, f"Template error: {exc}") from exc
    template_doc.save(str(target))
    return variables


__all__ = ["OPERATIONS", "EditState", "apply", "render_template", "validate"]
