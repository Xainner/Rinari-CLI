"""Inspección y lectura de presentaciones con python-pptx.

`inspect` da un mapa compacto (diapositivas, títulos, conteos, riesgos);
`read` devuelve el detalle de diapositivas concretas con identificadores
estables (`slide_id`, `shape_id`) para que una edición apunte a un objeto y no
a «el tercer texto que encuentre el parser». Lo que python-pptx no interpreta
(SmartArt, OLE, multimedia, animaciones) se declara, no se omite.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from rinari.documents.contracts import DocumentError, DocumentErrorCode

EMU_PER_PT = 12700
MAP_SLIDES = 300
READ_CHAR_BUDGET = 24_000
TEXT_PER_SHAPE = 4_000


def _pt(value: int | None) -> float | None:
    return round(value / EMU_PER_PT, 1) if value is not None else None


def _open(path: Path):
    from pptx import Presentation

    try:
        return Presentation(str(path))
    except Exception as exc:
        raise DocumentError(
            DocumentErrorCode.UNSUPPORTED_FORMAT, f"The presentation cannot be opened: {exc}"
        ) from exc


def _shape_kind(shape) -> str:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if shape.has_chart if hasattr(shape, "has_chart") else False:
        return "chart"
    if shape.has_table if hasattr(shape, "has_table") else False:
        return "table"
    shape_type = shape.shape_type
    if shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if shape_type in (MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT):
        return "ole"
    if shape_type == MSO_SHAPE_TYPE.MEDIA:
        return "media"
    if shape.is_placeholder:
        return "placeholder"
    xml = shape._element.xml if hasattr(shape, "_element") else ""
    if "drawingml/2006/diagram" in xml:
        return "smartart"
    if shape.has_text_frame:
        return "text"
    return "shape"


def _title(slide) -> str:
    shape = slide.shapes.title
    if shape is not None and shape.has_text_frame and shape.text_frame.text.strip():
        return shape.text_frame.text.strip()[:200]
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            return shape.text_frame.text.strip().splitlines()[0][:200]
    return ""


def _walk(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)


def inspect(path: Path) -> dict[str, Any]:
    presentation = _open(path)
    features = {
        k: 0
        for k in ("charts", "tables", "pictures", "groups", "ole", "media", "smartart", "notes")
    }
    slides: list[dict[str, Any]] = []
    total = len(presentation.slides)
    for index, slide in enumerate(presentation.slides, start=1):
        kinds = [_shape_kind(shape) for shape in _walk(slide.shapes)]
        for kind, key in (
            ("chart", "charts"),
            ("table", "tables"),
            ("picture", "pictures"),
            ("group", "groups"),
            ("ole", "ole"),
            ("media", "media"),
            ("smartart", "smartart"),
        ):
            features[key] += kinds.count(kind)
        has_notes = slide.has_notes_slide and bool(slide.notes_slide.notes_text_frame.text.strip())
        features["notes"] += int(has_notes)
        if index <= MAP_SLIDES:
            slides.append(
                {
                    "index": index,
                    "slide_id": slide.slide_id,
                    "layout": slide.slide_layout.name,
                    "title": _title(slide),
                    "shapes": len(kinds),
                    "notes": has_notes,
                    "hidden": slide._element.get("show") == "0",
                }
            )
    timing = sum(
        1
        for slide in presentation.slides
        if slide._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}timing")
        is not None
    )
    unsupported = []
    if features["smartart"]:
        unsupported.append({"feature": "smartart", "count": features["smartart"]})
    if features["ole"]:
        unsupported.append({"feature": "ole", "count": features["ole"]})
    if features["media"]:
        unsupported.append({"feature": "media", "count": features["media"]})
    if timing:
        unsupported.append({"feature": "animations", "count": timing})
    return {
        "kind": "pptx",
        "slide_count": total,
        "slide_size_pt": {
            "width": _pt(presentation.slide_width),
            "height": _pt(presentation.slide_height),
        },
        "layouts": [layout.name for layout in presentation.slide_layouts],
        "masters": len(presentation.slide_masters),
        "features": features,
        "unsupported": unsupported,
        "slides": slides,
        "truncated": total > MAP_SLIDES,
        "omitted_slides": max(0, total - MAP_SLIDES),
    }


def _text(shape) -> str | None:
    if not getattr(shape, "has_text_frame", False):
        return None
    return shape.text_frame.text[:TEXT_PER_SHAPE]


def _chart(shape) -> dict[str, Any]:
    chart = shape.chart
    plot = chart.plots[0] if len(chart.plots) else None
    data: dict[str, Any] = {"type": str(chart.chart_type).split(".")[-1].split(" ")[0]}
    if plot is not None:
        data["categories"] = [str(c) for c in plot.categories][:200]
        data["series"] = [
            {"name": series.name, "values": list(series.values)[:200]} for series in plot.series
        ][:20]
    data["has_title"] = chart.has_title
    if chart.has_title and chart.chart_title.has_text_frame:
        data["title"] = chart.chart_title.text_frame.text
    return data


def _table(shape) -> list[list[str]]:
    rows = []
    for row in list(shape.table.rows)[:100]:
        rows.append([cell.text[:500] for cell in list(row.cells)[:30]])
    return rows


def _shape_detail(shape) -> dict[str, Any]:
    kind = _shape_kind(shape)
    item: dict[str, Any] = {
        "shape_id": shape.shape_id,
        "name": shape.name,
        "kind": kind,
        "box_pt": {
            "x": _pt(shape.left),
            "y": _pt(shape.top),
            "w": _pt(shape.width),
            "h": _pt(shape.height),
        },
    }
    if shape.is_placeholder:
        item["placeholder"] = str(shape.placeholder_format.type).split(".")[-1].split(" ")[0]
    text = _text(shape)
    if text is not None:
        item["text"] = text
    if kind == "chart":
        item["chart"] = _chart(shape)
    elif kind == "table":
        item["table"] = _table(shape)
    elif kind == "picture":
        try:
            item["image"] = {
                "filename": shape.image.filename,
                "content_type": shape.image.content_type,
            }
        except Exception:
            item["image"] = {"linked": True}
    elif kind == "group":
        item["children"] = [_shape_detail(child) for child in shape.shapes]
    return item


def parse_selection(selection: str | None, count: int) -> list[int]:
    if not selection:
        return list(range(1, count + 1))
    numbers: list[int] = []
    for part in str(selection).split(","):
        part = part.strip()
        if not part:
            continue
        first, _, last = part.partition("-")
        try:
            start, end = int(first), int(last or first)
        except ValueError as exc:
            raise DocumentError(DocumentErrorCode.INVALID_SPEC, f"Bad selection {part!r}") from exc
        if start < 1 or end < start or end > count:
            raise DocumentError(
                DocumentErrorCode.INVALID_SPEC, f"Selection {part!r} is outside 1..{count}"
            )
        numbers.extend(range(start, end + 1))
    return sorted(set(numbers))


def read(path: Path, selection: str | None = None, cursor: int | None = None) -> dict[str, Any]:
    presentation = _open(path)
    slides = list(presentation.slides)
    wanted = parse_selection(selection, len(slides))
    if cursor:
        wanted = [n for n in wanted if n >= cursor]
    out: list[dict[str, Any]] = []
    used = 0
    next_cursor = None
    for number in wanted:
        slide = slides[number - 1]
        detail = {
            "index": number,
            "slide_id": slide.slide_id,
            "layout": slide.slide_layout.name,
            "shapes": [_shape_detail(shape) for shape in slide.shapes],
            "notes": slide.notes_slide.notes_text_frame.text[:TEXT_PER_SHAPE]
            if slide.has_notes_slide
            else "",
        }
        size = len(str(detail))
        if out and used + size > READ_CHAR_BUDGET:
            next_cursor = number
            break
        out.append(detail)
        used += size
    return {
        "kind": "pptx",
        "slides": out,
        "truncated": next_cursor is not None,
        "next_cursor": next_cursor,
        "pending_slides": [n for n in wanted if next_cursor and n >= next_cursor],
    }


def text_outline(path: Path, *, budget: int) -> tuple[str, bool]:
    """Texto por diapositiva para adjuntos (contexto inicial), con recorte explícito."""
    presentation = _open(path)
    parts: list[str] = []
    used = 0
    for index, slide in enumerate(presentation.slides, start=1):
        lines = [f"[Slide {index}]"]
        for shape in _walk(slide.shapes):
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            elif getattr(shape, "has_table", False) and shape.has_table:
                lines.extend("\t".join(row) for row in _table(shape))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            lines.append("[Notes] " + slide.notes_slide.notes_text_frame.text.strip())
        block = "\n".join(lines)
        if used + len(block) > budget:
            parts.append(f"[Slides {index}-{len(presentation.slides)} not included: limit]")
            return "\n\n".join(parts), True
        parts.append(block)
        used += len(block)
    return "\n\n".join(parts), False
