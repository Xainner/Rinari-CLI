"""Compilador de DeckSpec a PPTX nativo con python-pptx.

La skill decide la narrativa; aquí se calcula la geometría. Texto, tablas y
gráficos son objetos nativos y editables (los gráficos llevan sus datos
incrustados). Cada forma se nombra `rinari:<slide>:<rol>` para que una
edición posterior la encuentre por su identidad, no por su orden.

Lo que no cabe no se encoge hasta ser ilegible: se registra como hallazgo
(`overflow`) con la acción posible (resumir, dividir, pasar al anexo).
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from rinari.documents.contracts import DocumentError, DocumentErrorCode
from rinari.documents.design import measure
from rinari.documents.design.tokens import Theme, theme

SLIDE_W = 960.0
SLIDE_H = 540.0
LAYOUTS = (
    "cover",
    "section",
    "statement",
    "bullets",
    "summary",
    "kpi",
    "chart",
    "comparison",
    "matrix",
    "table",
    "timeline",
    "process",
    "image",
    "quote",
    "closing",
)
CHART_TYPES = (
    "bar",
    "column",
    "stacked_bar",
    "stacked_column",
    "line",
    "area",
    "pie",
    "doughnut",
)
MAX_SLIDES = 120


@dataclass(slots=True)
class Element:
    role: str
    box: tuple[float, float, float, float]
    text: str = ""
    size: float | None = None
    min_size: float | None = None
    color: str | None = None
    background: str | None = None
    fits: bool = True
    overlap_ok: bool = False

    def to_dict(self) -> dict[str, Any]:
        return {
            "role": self.role,
            "box_pt": [round(v, 1) for v in self.box],
            "size": self.size,
            "fits": self.fits,
            "chars": len(self.text),
        }


@dataclass(slots=True)
class SlidePlan:
    id: str
    layout: str
    elements: list[Element] = field(default_factory=list)
    findings: list[dict[str, Any]] = field(default_factory=list)


def _pt(value: float) -> int:
    return round(value * 12700)


def _rgb(value: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(value.lstrip("#").upper())


class _Builder:
    def __init__(self, spec: dict[str, Any], resources: dict[str, str], prs=None) -> None:
        from pptx import Presentation

        self.spec = spec
        self.theme: Theme = theme(spec.get("theme"))
        self.resources = resources
        if prs is None:
            self.prs = Presentation()
            ratio = spec.get("aspect_ratio", "16:9")
            self.width = SLIDE_W if ratio == "16:9" else 720.0
            self.height = SLIDE_H
            self.prs.slide_width = _pt(self.width)
            self.prs.slide_height = _pt(self.height)
        else:
            # Una diapositiva nueva en un deck ajeno usa su lienzo, no el nuestro.
            self.prs = prs
            self.width = prs.slide_width / 12700
            self.height = prs.slide_height / 12700
        self.blank = _blank_layout(self.prs)
        self.plans: list[SlidePlan] = []

    # -- primitivas ------------------------------------------------------------------
    def _text(
        self,
        slide,
        plan: SlidePlan,
        role: str,
        text: str,
        box: tuple[float, float, float, float],
        *,
        size: float,
        min_size: float | None = None,
        bold: bool = False,
        italic: bool = False,
        color: str | None = None,
        font: str | None = None,
        align: str = "left",
        anchor: str = "top",
        bullets: list[str] | None = None,
        spacing: float = 1.1,
        background: str | None = None,
    ):
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

        x, y, w, h = box
        family = font or self.theme.body_font
        content = "\n".join(bullets) if bullets is not None else str(text)
        inner_w = w - 8
        fitted, fits = measure.fit(
            content if bullets is None else "\n".join("• " + b for b in bullets),
            inner_w,
            h - 6,
            family,
            size,
            min_size or size,
            bold=bold,
            spacing=spacing,
        )
        shape = slide.shapes.add_textbox(_pt(x), _pt(y), _pt(w), _pt(h))
        shape.name = f"rinari:{plan.id}:{role}"
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = _pt(4)
        frame.margin_top = frame.margin_bottom = _pt(3)
        frame.vertical_anchor = {
            "top": MSO_ANCHOR.TOP,
            "middle": MSO_ANCHOR.MIDDLE,
            "bottom": MSO_ANCHOR.BOTTOM,
        }[anchor]
        lines = bullets if bullets is not None else str(text).split("\n")
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = {
                "left": PP_ALIGN.LEFT,
                "center": PP_ALIGN.CENTER,
                "right": PP_ALIGN.RIGHT,
            }[align]
            paragraph.line_spacing = spacing
            if bullets is not None:
                _bullet(paragraph, self.theme.accent)
                paragraph.space_after = _pt(fitted * 0.45)
            run = paragraph.add_run()
            run.text = line
            run.font.size = _pt(fitted)
            run.font.bold = bold
            run.font.italic = italic
            run.font.name = family
            run.font.color.rgb = _rgb(color or self.theme.text)
        element = Element(
            role=role,
            box=box,
            text=content,
            size=fitted,
            min_size=min_size or size,
            color=color or self.theme.text,
            background=background or self.theme.background,
            fits=fits,
        )
        plan.elements.append(element)
        if not fits:
            plan.findings.append(
                {
                    "code": "TEXT_OVERFLOW",
                    "severity": "error",
                    "slide": plan.id,
                    "role": role,
                    "detail": f"{role} does not fit at {min_size or size} pt",
                    "action": "Summarize it, split the slide or move detail to the appendix",
                }
            )
        return shape

    def _height(
        self,
        text: str,
        width: float,
        size: float,
        *,
        bold: bool = False,
        font: str | None = None,
        spacing: float = 1.1,
    ) -> float:
        needed, _ = measure.text_height(
            text, width - 8, font or self.theme.body_font, size, bold=bold, spacing=spacing
        )
        return needed + 8

    def _rect(
        self,
        slide,
        plan: SlidePlan,
        role: str,
        box: tuple[float, float, float, float],
        fill: str,
        *,
        rounded: bool = False,
        line: str | None = None,
        shape_kind=None,
    ):
        from pptx.enum.shapes import MSO_SHAPE

        x, y, w, h = box
        kind = shape_kind or (MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE)
        shape = slide.shapes.add_shape(kind, _pt(x), _pt(y), _pt(w), _pt(h))
        shape.name = f"rinari:{plan.id}:{role}"
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
        if line:
            shape.line.color.rgb = _rgb(line)
            shape.line.width = _pt(0.75)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        if rounded and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            radius = self.theme.radius / max(1.0, min(w, h))
            shape.adjustments[0] = min(0.5, radius)
        plan.elements.append(Element(role=role, box=box, overlap_ok=True))
        return shape

    def _background(self, slide, color: str | None = None) -> None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(color or self.theme.background)

    # -- piezas comunes ------------------------------------------------------------
    def _title(self, slide, plan: SlidePlan, text: str, *, top: float | None = None) -> float:
        t = self.theme
        y = top if top is not None else t.margin_top
        self._rect(slide, plan, "accent", (t.margin_x, y, 36, 4), t.accent)
        self._text(
            slide,
            plan,
            "title",
            text,
            (t.margin_x, y + 12, self.width - 2 * t.margin_x, 84),
            size=t.heading_size,
            min_size=22,
            bold=True,
            font=t.heading_font,
            spacing=0.95,
        )
        return y + 12 + 84 + 10

    def _footer(self, slide, plan: SlidePlan, index: int, source: str | None) -> None:
        t = self.theme
        y = self.height - t.margin_bottom + 4
        if source:
            self._text(
                slide,
                plan,
                "source",
                source,
                (t.margin_x, y, self.width - 2 * t.margin_x - 60, 24),
                size=10,
                min_size=10,
                color=t.muted,
            )
        number = self._text(
            slide,
            plan,
            "number",
            str(index),
            (self.width - t.margin_x - 40, y, 40, 24),
            size=10,
            color=t.muted,
            align="right",
        )
        _slide_number_field(number.text_frame.paragraphs[0])

    def _content_box(self, top: float) -> tuple[float, float, float, float]:
        t = self.theme
        bottom = self.height - t.margin_bottom - 8
        return t.margin_x, top, self.width - 2 * t.margin_x, bottom - top

    # -- layouts -----------------------------------------------------------------------
    def build(self) -> bytes:
        slides = self.spec.get("slides")
        if not isinstance(slides, list) or not slides:
            raise DocumentError(DocumentErrorCode.INVALID_SPEC, "slides must be a non-empty list")
        if len(slides) > MAX_SLIDES:
            raise DocumentError(
                DocumentErrorCode.DOCUMENT_LIMIT_EXCEEDED, f"At most {MAX_SLIDES} slides"
            )
        unknown = sorted(set(self.spec) - DECK_KEYS)
        if unknown:
            raise DocumentError(
                DocumentErrorCode.INVALID_SPEC,
                f"Unknown deck fields: {', '.join(unknown)}",
                details={"allowed": sorted(DECK_KEYS)},
            )
        seen: set[str] = set()
        for index, spec in enumerate(slides, start=1):
            slide_id = str(spec.get("id") or f"s{index:02d}") if isinstance(spec, dict) else ""
            if slide_id in seen:
                raise DocumentError(
                    DocumentErrorCode.INVALID_SPEC, f"Duplicate slide id {slide_id}"
                )
            seen.add(slide_id)
            self.add(spec, index, slide_id)
        props = self.prs.core_properties
        props.title = str(self.spec.get("title") or "")[:250]
        props.language = str(self.spec.get("language") or "es")
        props.author = str(self.spec.get("author") or "")[:120]
        props.comments = "Generated by Rinari"
        buffer = io.BytesIO()
        self.prs.save(buffer)
        return buffer.getvalue()

    def add(self, spec: Any, index: int, slide_id: str):
        """Compone una diapositiva al final del deck y la devuelve."""
        validate_slide(spec, index)
        layout = spec.get("layout", "bullets")
        slide = self.prs.slides.add_slide(self.blank)
        for placeholder in list(slide.placeholders):
            placeholder._element.getparent().remove(placeholder._element)
        plan = SlidePlan(id=slide_id, layout=layout)
        self._background(slide, self.theme.accent if layout == "section" else None)
        getattr(self, f"_layout_{layout}")(slide, plan, spec)
        if layout not in ("cover", "section", "closing"):
            self._footer(slide, plan, index, spec.get("source"))
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(spec["notes"])
        self.plans.append(plan)
        return slide

    def _layout_cover(self, slide, plan, spec):
        t = self.theme
        x = t.margin_x + 8
        if spec.get("eyebrow"):
            self._text(
                slide,
                plan,
                "eyebrow",
                spec["eyebrow"],
                (x, 150, 700, 26),
                size=14,
                bold=True,
                color=t.ink,
            )
        self._rect(slide, plan, "accent", (x, 186, 56, 5), t.accent)
        title = _required(spec, "title")
        width = self.width - 2 * x
        title_h = min(
            150.0,
            self._height(
                title, width, t.title_size + 6, bold=True, font=t.heading_font, spacing=0.95
            ),
        )
        self._text(
            slide,
            plan,
            "title",
            title,
            (x, 200, width, max(title_h, 60)),
            size=t.title_size + 6,
            min_size=t.min_title_size,
            bold=True,
            font=t.heading_font,
            spacing=0.95,
        )
        if spec.get("subtitle"):
            self._text(
                slide,
                plan,
                "subtitle",
                spec["subtitle"],
                (x, 200 + max(title_h, 60) + 14, width, 60),
                size=20,
                min_size=16,
                color=t.muted,
            )
        if spec.get("meta"):
            self._text(
                slide,
                plan,
                "meta",
                spec["meta"],
                (x, self.height - 70, 600, 24),
                size=12,
                color=t.muted,
            )

    def _layout_section(self, slide, plan, spec):
        t = self.theme
        x = t.margin_x + 8
        if spec.get("number"):
            self._text(
                slide,
                plan,
                "number",
                str(spec["number"]),
                (x, 170, 200, 60),
                size=44,
                bold=True,
                color=t.accent_text,
            )
        self._text(
            slide,
            plan,
            "title",
            _required(spec, "title"),
            (x, 236, self.width - 2 * x, 110),
            size=t.title_size + 4,
            min_size=t.min_title_size,
            bold=True,
            font=t.heading_font,
            color=t.accent_text,
            spacing=0.95,
            background=t.accent,
        )
        if spec.get("subtitle"):
            self._text(
                slide,
                plan,
                "subtitle",
                spec["subtitle"],
                (x, 350, self.width - 2 * x, 50),
                size=18,
                color=t.accent_text,
                background=t.accent,
            )

    def _layout_statement(self, slide, plan, spec):
        t = self.theme
        self._rect(slide, plan, "accent", (t.margin_x, 150, 6, 170), t.accent)
        self._text(
            slide,
            plan,
            "title",
            _required(spec, "title"),
            (t.margin_x + 28, 140, self.width - 2 * t.margin_x - 28, 190),
            size=34,
            min_size=26,
            bold=True,
            font=t.heading_font,
            anchor="middle",
            spacing=1.0,
        )
        if spec.get("support"):
            self._text(
                slide,
                plan,
                "support",
                spec["support"],
                (t.margin_x + 28, 340, self.width - 2 * t.margin_x - 28, 90),
                size=18,
                min_size=14,
                color=t.muted,
            )

    def _layout_bullets(self, slide, plan, spec):
        top = self._title(slide, plan, _required(spec, "title"))
        x, y, w, h = self._content_box(top)
        points = _strings(spec, "points", 1, 8)
        self._text(
            slide,
            plan,
            "body",
            "",
            (x, y, w, h),
            size=self.theme.body_size + 2,
            min_size=self.theme.min_body_size,
            bullets=points,
            spacing=1.1,
        )

    def _layout_summary(self, slide, plan, spec):
        t = self.theme
        top = self._title(slide, plan, _required(spec, "title"))
        x, y, w, h = self._content_box(top)
        points = spec.get("points")
        if not isinstance(points, list) or not 2 <= len(points) <= 6:
            raise DocumentError(
                DocumentErrorCode.INVALID_SPEC, f"{plan.id}: summary needs 2-6 points"
            )
        rows = len(points)
        row_h = min(92.0, h / rows)
        head_w = min(260.0, (w - 44) * 0.32)
        for index, point in enumerate(points):
            head = point.get("head") if isinstance(point, dict) else None
            text = point.get("text") if isinstance(point, dict) else str(point)
            ry = y + index * row_h
            if index:
                self._rect(slide, plan, f"rule{index}", (x + 44, ry - 1, w - 44, 0.75), t.border)
            self._text(
                slide,
                plan,
                f"marker{index + 1}",
                str(index + 1),
                (x, ry + 6, 34, row_h - 14),
                size=22,
                bold=True,
                color=t.ink,
                font=t.heading_font,
            )
            if head:
                self._text(
                    slide,
                    plan,
                    f"head{index + 1}",
                    head,
                    (x + 44, ry + 8, head_w, row_h - 16),
                    size=19,
                    min_size=14,
                    bold=True,
                )
                self._text(
                    slide,
                    plan,
                    f"text{index + 1}",
                    text or "",
                    (x + 44 + head_w + 20, ry + 8, w - 64 - head_w, row_h - 16),
                    size=17,
                    min_size=13,
                    color=t.muted,
                )
            else:
                self._text(
                    slide,
                    plan,
                    f"text{index + 1}",
                    text or "",
                    (x + 44, ry + 8, w - 44, row_h - 16),
                    size=19,
                    min_size=14,
                )

    def _layout_kpi(self, slide, plan, spec):
        t = self.theme
        top = self._title(slide, plan, _required(spec, "title"))
        x, y, w, h = self._content_box(top)
        kpis = spec.get("kpis")
        if not isinstance(kpis, list) or not 1 <= len(kpis) <= 4:
            raise DocumentError(DocumentErrorCode.INVALID_SPEC, f"{plan.id}: kpi needs 1-4 kpis")
        gap = t.gutter
        card_w = (w - gap * (len(kpis) - 1)) / len(kpis)
        card_h = min(190.0, h - (60 if spec.get("takeaway") else 0))
        for index, kpi in enumerate(kpis):
            cx = x + index * (card_w + gap)
            self._rect(
                slide, plan, f"card{index + 1}", (cx, y, card_w, card_h), t.surface, rounded=True
            )
            self._text(
                slide,
                plan,
                f"label{index + 1}",
                str(kpi.get("label", "")),
                (cx + 18, y + 16, card_w - 36, 40),
                size=14,
                min_size=11,
                color=t.muted,
                background=t.surface,
            )
            self._text(
                slide,
                plan,
                f"value{index + 1}",
                str(_required(kpi, "value")),
                (cx + 18, y + 58, card_w - 36, 70),
                size=40,
                min_size=26,
                bold=True,
                font=t.heading_font,
                background=t.surface,
            )
            delta = kpi.get("delta")
            if delta not in (None, ""):
                trend = kpi.get("trend") or (
                    "up"
                    if str(delta).startswith("+")
                    else "down"
                    if str(delta).startswith(("-", "\u2212"))
                    else "flat"
                )
                good = kpi.get("good", "up")
                color = (
                    t.muted if trend == "flat" else (t.positive if trend == good else t.negative)
                )
                arrow = {"up": "▲ ", "down": "▼ ", "flat": "■ "}.get(trend, "")
                self._text(
                    slide,
                    plan,
                    f"delta{index + 1}",
                    f"{arrow}{delta}",
                    (cx + 18, y + 128, card_w - 36, 28),
                    size=14,
                    bold=True,
                    color=color,
                    background=t.surface,
                )
            if kpi.get("note"):
                self._text(
                    slide,
                    plan,
                    f"note{index + 1}",
                    str(kpi["note"]),
                    (cx + 18, y + 156, card_w - 36, card_h - 160),
                    size=11,
                    min_size=10,
                    color=t.muted,
                    background=t.surface,
                )
        if spec.get("takeaway"):
            self._text(
                slide,
                plan,
                "takeaway",
                spec["takeaway"],
                (x, y + card_h + 18, w, 44),
                size=16,
                min_size=13,
                color=t.text,
            )

    def _layout_chart(self, slide, plan, spec):
        t = self.theme
        top = self._title(slide, plan, _required(spec, "title"))
        x, y, w, h = self._content_box(top)
        takeaway = spec.get("takeaway")
        chart_w = w * 0.66 if takeaway else w
        self._chart(slide, plan, spec.get("chart") or {}, (x, y, chart_w, h))
        if takeaway:
            panel_x = x + chart_w + t.gutter
            panel_w = w - chart_w - t.gutter
            self._rect(slide, plan, "panel", (panel_x, y, panel_w, h), t.surface, rounded=True)
            self._rect(slide, plan, "panel-accent", (panel_x + 18, y + 22, 28, 4), t.accent)
            self._text(
                slide,
                plan,
                "takeaway",
                takeaway,
                (panel_x + 18, y + 36, panel_w - 36, h - 54),
                size=18,
                min_size=13,
                font=t.heading_font,
                background=t.surface,
            )

    def _chart(self, slide, plan, chart: dict[str, Any], box):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

        t = self.theme
        kind = chart.get("type", "column")
        if kind not in CHART_TYPES:
            raise DocumentError(
                DocumentErrorCode.INVALID_SPEC,
                f"{plan.id}: unknown chart type {kind!r}",
                details={"types": list(CHART_TYPES)},
            )
        categories = chart.get("categories")
        series = chart.get("series")
        if (
            not isinstance(categories, list)
            or not categories
            or not isinstance(series, list)
            or not series
        ):
            raise DocumentError(
                DocumentErrorCode.INVALID_SPEC, f"{plan.id}: chart needs categories and series"
            )
        data = CategoryChartData(number_format=chart.get("number_format") or "General")
        data.categories = [str(c) for c in categories]
        for item in series:
            values = item.get("values") if isinstance(item, dict) else None
            if not isinstance(values, list) or len(values) != len(categories):
                raise DocumentError(
                    DocumentErrorCode.INVALID_SPEC,
                    f"{plan.id}: series {_series_name(item)!r} needs one value per category",
                )
            if not all(
                v is None or (isinstance(v, (int, float)) and not isinstance(v, bool))
                for v in values
            ):
                raise DocumentError(
                    DocumentErrorCode.INVALID_SPEC, f"{plan.id}: chart values must be numbers"
                )
            data.add_series(str(item.get("name") or ""), values)
        xl = {
            "bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
            "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
            "line": XL_CHART_TYPE.LINE_MARKERS,
            "area": XL_CHART_TYPE.AREA,
            "pie": XL_CHART_TYPE.PIE,
            "doughnut": XL_CHART_TYPE.DOUGHNUT,
        }[kind]
        x, y, w, h = box
        frame = slide.shapes.add_chart(xl, _pt(x), _pt(y), _pt(w), _pt(h), data)
        frame.name = f"rinari:{plan.id}:chart"
        plot_chart = frame.chart
        plot_chart.has_title = False
        plot_chart.font.size = _pt(12)
        plot_chart.font.name = t.body_font
        plot_chart.font.color.rgb = _rgb(t.muted)
        circular = kind in ("pie", "doughnut")
        many = len(series) > 1
        plot_chart.has_legend = many or circular
        if plot_chart.has_legend:
            plot_chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            plot_chart.legend.include_in_layout = False
            plot_chart.legend.font.size = _pt(12)
        plot = plot_chart.plots[0]
        if circular:
            for index, point in enumerate(plot.series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = _rgb(t.palette[index % len(t.palette)])
        else:
            for index, item in enumerate(plot.series):
                color = _rgb(t.palette[index % len(t.palette)])
                if kind == "line":
                    item.format.line.color.rgb = color
                    item.format.line.width = _pt(2.25)
                    item.smooth = False
                    # El marcador del color de su serie, no el de la paleta de Office.
                    from pptx.enum.chart import XL_MARKER_STYLE

                    item.marker.style = XL_MARKER_STYLE.CIRCLE
                    item.marker.size = 7
                    item.marker.format.fill.solid()
                    item.marker.format.fill.fore_color.rgb = color
                    item.marker.format.line.color.rgb = color
                else:
                    item.format.fill.solid()
                    item.format.fill.fore_color.rgb = color
                    item.invert_if_negative = False
            if hasattr(plot, "gap_width"):
                plot.gap_width = 60
            value_axis = plot_chart.value_axis
            value_axis.has_major_gridlines = True
            value_axis.major_gridlines.format.line.color.rgb = _rgb(t.border)
            value_axis.format.line.fill.background()
            value_axis.tick_labels.font.size = _pt(11)
            category_axis = plot_chart.category_axis
            category_axis.format.line.color.rgb = _rgb(t.border)
            category_axis.tick_labels.font.size = _pt(11)
            # Con valores negativos las etiquetas van al borde, no sobre la barra.
            from pptx.enum.chart import XL_TICK_LABEL_POSITION

            category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        small = len(categories) <= 8 and (not many or circular)
        if small or chart.get("data_labels"):
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.font.size = _pt(11)
            labels.font.color.rgb = _rgb(t.text)
            labels.number_format = chart.get("number_format") or "General"
            labels.number_format_is_linked = not chart.get("number_format")
        plan.elements.append(Element(role="chart", box=box))

    def _layout_comparison(self, slide, plan, spec):
        t = self.theme
        top = self._title(slide, plan, _required(spec, "title"))
        x, y, w, h = self._content_box(top)
        columns = spec.get("columns")
        if not isinstance(columns, list) or not 2 <= len(columns) <= 3:
            raise DocumentError(
                DocumentErrorCode.INVALID_SPEC, f"{plan.id}: comparison needs 2-3 columns"
            )
        gap = t.gutter
        col_w = (w - gap * (len(columns) - 1)) / len(columns)
        lists = [_strings(column, "points", 1, 6) for column in columns]
        needed = max(
            self._height("\n".join("• " + p for p in points), col_w - 54, 17) * 1.25
            for points in lists
        )
        card_h = min(h, max(180.0, 74 + needed + 24))
        for index, column in enumerate(columns):
            cx = x + index * (col_w + gap)
            self._rect(
                slide, plan, f"card{index + 1}", (cx, y, col_w, card_h), t.surface, rounded=True
            )
            self._rect(
                slide, plan, f"bar{index + 1}", (cx, y, col_w, 5), t.palette[index % len(t.palette)]
            )
            self._text(
                slide,
                plan,
                f"heading{index + 1}",
                str(_required(column, "heading")),
                (cx + 18, y + 20, col_w - 36, 44),
                size=21,
                min_size=15,
                bold=True,
                font=t.heading_font,
                background=t.surface,
            )
            self._text(
                slide,
                plan,
                f"points{index + 1}",
                "",
                (cx + 18, y + 72, col_w - 36, card_h - 90),
                size=17,
                min_size=12,
                bullets=lists[index],
                background=t.surface,
            )

    def _layout_matrix(self, slide, plan, spec):
        t = self.theme
        top = self._title(slide, plan, _required(spec, "title"))
        x, y, w, h = self._content_box(top)
        quadrants = spec.get("quadrants")
        if not isinstance(quadrants, list) or len(quadrants) != 4:
            raise DocumentError(
                DocumentErrorCode.INVALID_SPEC, f"{plan.id}: matrix needs 4 quadrants"
            )
        label = 26 if spec.get("y_axis") else 0
        bottom = 24 if spec.get("x_axis") else 0
        gx, gw, gh = x + label, w - label, h - bottom
        cell_w, cell_h = (gw - 12) / 2, (gh - 12) / 2
        for index, quadrant in enumerate(quadrants):
            col, row = index % 2, index // 2
            cx, cy = gx + col * (cell_w + 12), y + row * (cell_h + 12)
            emphasis = bool(quadrant.get("highlight"))
            fill = t.accent if emphasis else t.surface
            ink = t.accent_text if emphasis else t.text
            self._rect(
                slide, plan, f"cell{index + 1}", (cx, cy, cell_w, cell_h), fill, rounded=True
            )
            self._text(
                slide,
                plan,
                f"heading{index + 1}",
                str(_required(quadrant, "heading")),
                (cx + 16, cy + 12, cell_w - 32, 34),
                size=17,
                min_size=13,
                bold=True,
                color=ink,
                background=fill,
            )
            if quadrant.get("text"):
                self._text(
                    slide,
                    plan,
                    f"text{index + 1}",
                    str(quadrant["text"]),
                    (cx + 16, cy + 48, cell_w - 32, cell_h - 58),
                    size=14,
                    min_size=11,
                    color=ink if emphasis else t.muted,
                    background=fill,
                )
        if spec.get("x_axis"):
            self._text(
                slide,
                plan,
                "x-axis",
                spec["x_axis"],
                (gx, y + gh + 4, gw, 20),
                size=11,
                color=t.muted,
                align="center",
            )
        if spec.get("y_axis"):
            box = self._text(
                slide,
                plan,
                "y-axis",
                spec["y_axis"],
                (x - gh / 2 + 10, y + gh / 2 - 10, gh, 20),
                size=11,
                color=t.muted,
                align="center",
            )
            box.rotation = -90
            plan.elements[-1].overlap_ok = True

    def _layout_table(self, slide, plan, spec):
        from pptx.enum.text import PP_ALIGN

        t = self.theme
        top = self._title(slide, plan, _required(spec, "title"))
        x, y, w, h = self._content_box(top)
        columns = spec.get("columns")
        rows = spec.get("rows")
        if not isinstance(columns, list) or not columns or not isinstance(rows, list):
            raise DocumentError(
                DocumentErrorCode.INVALID_SPEC, f"{plan.id}: table needs columns and rows"
            )
        if len(rows) > 14 or len(columns) > 8:
            plan.findings.append(
                {
                    "code": "TABLE_TOO_DENSE",
                    "severity": "error",
                    "slide": plan.id,
                    "detail": f"{len(rows)} rows x {len(columns)} columns",
                    "action": "Split the table or move it to the appendix",
                }
            )
        n_rows = len(rows) + 1
        row_h = min(34.0, h / max(1, n_rows))
        size = 14 if row_h >= 30 else 12 if row_h >= 24 else 10
        frame = slide.shapes.add_table(
            n_rows, len(columns), _pt(x), _pt(y), _pt(w), _pt(row_h * n_rows)
        )
        frame.name = f"rinari:{plan.id}:table"
        table = frame.table
        numeric = [
            all(_is_number(r[c]) for r in rows if c < len(r) and r[c] not in (None, "")) and rows
            for c in range(len(columns))
        ]
        language = str(self.spec.get("language") or "es")
        decimals = [
            max(
                (_decimals(r[c]) for r in rows if c < len(r) and isinstance(r[c], float)), default=0
            )
            for c in range(len(columns))
        ]
        highlight = spec.get("highlight_row")
        for c, heading in enumerate(columns):
            cell = table.cell(0, c)
            _cell(
                cell,
                str(heading),
                size,
                True,
                t.accent_text,
                t.accent,
                PP_ALIGN.RIGHT if numeric[c] else PP_ALIGN.LEFT,
                t.body_font,
                t.accent,
            )
        for r, row in enumerate(rows, start=1):
            if not isinstance(row, list) or len(row) != len(columns):
                raise DocumentError(
                    DocumentErrorCode.INVALID_SPEC, f"{plan.id}: row {r} needs {len(columns)} cells"
                )
            fill = t.surface if r % 2 == 0 else t.background
            bold = highlight is not None and r - 1 == highlight
            for c, value in enumerate(row):
                _cell(
                    table.cell(r, c),
                    _format(value, language, decimals[c]),
                    size,
                    bold,
                    t.text,
                    fill,
                    PP_ALIGN.RIGHT if numeric[c] else PP_ALIGN.LEFT,
                    t.body_font,
                    t.border,
                )
        for r in range(n_rows):
            table.rows[r].height = _pt(row_h)
        plan.elements.append(
            Element(
                role="table",
                box=(x, y, w, row_h * n_rows),
                size=size,
                min_size=10,
                fits=row_h * n_rows <= h + 1,
            )
        )

    def _layout_timeline(self, slide, plan, spec):
        t = self.theme
        top = self._title(slide, plan, _required(spec, "title"))
        x, y, w, h = self._content_box(top)
        events = spec.get("events")
        if not isinstance(events, list) or not 2 <= len(events) <= 6:
            raise DocumentError(
                DocumentErrorCode.INVALID_SPEC, f"{plan.id}: timeline needs 2-6 events"
            )
        line_y = y + 70
        self._rect(slide, plan, "line", (x, line_y, w, 3), t.border)
        step = w / len(events)
        for index, event in enumerate(events):
            cx = x + index * step
            self._rect(
                slide,
                plan,
                f"dot{index + 1}",
                (cx, line_y - 6, 15, 15),
                t.palette[index % len(t.palette)],
                rounded=True,
            )
            self._text(
                slide,
                plan,
                f"date{index + 1}",
                str(_required(event, "date")),
                (cx, y + 20, step - 16, 34),
                size=15,
                min_size=12,
                bold=True,
                color=t.ink,
            )
            self._text(
                slide,
                plan,
                f"heading{index + 1}",
                str(_required(event, "heading")),
                (cx, line_y + 26, step - 16, 50),
                size=16,
                min_size=12,
                bold=True,
            )
            if event.get("text"):
                self._text(
                    slide,
                    plan,
                    f"text{index + 1}",
                    str(event["text"]),
                    (cx, line_y + 78, step - 16, h - (line_y - y) - 86),
                    size=13,
                    min_size=10,
                    color=t.muted,
                )

    def _layout_process(self, slide, plan, spec):
        t = self.theme
        top = self._title(slide, plan, _required(spec, "title"))
        x, y, w, h = self._content_box(top)
        steps = spec.get("steps")
        if not isinstance(steps, list) or not 2 <= len(steps) <= 6:
            raise DocumentError(
                DocumentErrorCode.INVALID_SPEC, f"{plan.id}: process needs 2-6 steps"
            )
        gap = 22
        step_w = (w - gap * (len(steps) - 1)) / len(steps)
        tallest = max(self._height(str(s.get("text") or ""), step_w - 32, 15) for s in steps)
        card_h = min(h - 24, max(170.0, 104 + tallest + 26))
        for index, step in enumerate(steps):
            cx = x + index * (step_w + gap)
            self._rect(
                slide,
                plan,
                f"card{index + 1}",
                (cx, y + 24, step_w, card_h),
                t.surface,
                rounded=True,
            )
            self._rect(
                slide, plan, f"badge{index + 1}", (cx + 16, y + 6, 38, 38), t.accent, rounded=True
            )
            self._text(
                slide,
                plan,
                f"n{index + 1}",
                str(index + 1),
                (cx + 16, y + 10, 38, 30),
                size=17,
                bold=True,
                color=t.accent_text,
                align="center",
                background=t.accent,
            )
            self._text(
                slide,
                plan,
                f"heading{index + 1}",
                str(_required(step, "heading")),
                (cx + 16, y + 60, step_w - 32, 50),
                size=19,
                min_size=13,
                bold=True,
                background=t.surface,
            )
            if step.get("text"):
                self._text(
                    slide,
                    plan,
                    f"text{index + 1}",
                    str(step["text"]),
                    (cx + 16, y + 110, step_w - 32, card_h - 100),
                    size=15,
                    min_size=10,
                    color=t.muted,
                    background=t.surface,
                )

    def _layout_image(self, slide, plan, spec):
        from PIL import Image

        t = self.theme
        top = self._title(slide, plan, _required(spec, "title"))
        x, y, w, h = self._content_box(top)
        ref = _required(spec, "image")
        path = self.resources.get(str(ref))
        if path is None:
            raise DocumentError(
                DocumentErrorCode.NOT_FOUND, f"{plan.id}: image {ref!r} was not provided"
            )
        caption = spec.get("caption")
        avail_h = h - (34 if caption else 0)
        with Image.open(path) as image:
            iw, ih = image.size
        scale = min(w / iw, avail_h / ih)
        pw, ph = iw * scale, ih * scale
        px, py = x + (w - pw) / 2, y + (avail_h - ph) / 2
        picture = slide.shapes.add_picture(path, _pt(px), _pt(py), _pt(pw), _pt(ph))
        picture.name = f"rinari:{plan.id}:image"
        plan.elements.append(Element(role="image", box=(px, py, pw, ph)))
        if caption:
            self._text(
                slide,
                plan,
                "caption",
                caption,
                (x, y + avail_h + 6, w, 26),
                size=12,
                min_size=10,
                color=t.muted,
                align="center",
            )

    def _layout_quote(self, slide, plan, spec):
        t = self.theme
        self._text(
            slide,
            plan,
            "mark",
            "“",
            (t.margin_x, 70, 120, 120),
            size=110,
            bold=True,
            color=t.ink,
            font="Georgia",
        )
        self._text(
            slide,
            plan,
            "quote",
            _required(spec, "quote"),
            (t.margin_x + 40, 150, self.width - 2 * t.margin_x - 80, 210),
            size=28,
            min_size=18,
            italic=True,
            font=t.heading_font,
            spacing=1.1,
        )
        attribution = " · ".join(str(v) for v in (spec.get("author"), spec.get("role")) if v)
        if attribution:
            self._text(
                slide,
                plan,
                "author",
                attribution,
                (t.margin_x + 40, 380, self.width - 2 * t.margin_x - 80, 30),
                size=15,
                bold=True,
                color=t.muted,
            )

    def _layout_closing(self, slide, plan, spec):
        t = self.theme
        x = t.margin_x + 8
        self._rect(slide, plan, "accent", (x, 200, 56, 5), t.accent)
        title = _required(spec, "title")
        width = self.width - 2 * x
        title_h = max(
            60.0,
            min(
                110.0, self._height(title, width, t.title_size + 4, bold=True, font=t.heading_font)
            ),
        )
        self._text(
            slide,
            plan,
            "title",
            title,
            (x, 214, width, title_h),
            size=t.title_size + 4,
            min_size=t.min_title_size,
            bold=True,
            font=t.heading_font,
        )
        if spec.get("subtitle"):
            self._text(
                slide,
                plan,
                "subtitle",
                spec["subtitle"],
                (x, 214 + title_h + 12, width, 50),
                size=19,
                min_size=15,
                color=t.muted,
            )
        if spec.get("contact"):
            self._text(
                slide,
                plan,
                "contact",
                spec["contact"],
                (x, self.height - 80, 700, 26),
                size=13,
                color=t.muted,
            )


def _slide_number_field(paragraph) -> None:
    """El número es un campo de PowerPoint: sigue siendo cierto al reordenar."""
    import uuid

    from pptx.oxml.ns import qn

    run = paragraph.runs[0]._r
    field = run.makeelement(qn("a:fld"), {"id": "{" + str(uuid.uuid4()).upper() + "}"})
    field.set("type", "slidenum")
    for child in list(run):
        field.append(child)
    run.addprevious(field)
    run.getparent().remove(run)


def _bullet(paragraph, color: str) -> None:
    from pptx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("marL", str(_pt(18)))
    p_pr.set("indent", str(-_pt(18)))
    for tag in ("a:buNone", "a:buChar", "a:buClr"):
        for node in p_pr.findall(qn(tag)):
            p_pr.remove(node)
    clr = p_pr.makeelement(qn("a:buClr"), {})
    srgb = clr.makeelement(qn("a:srgbClr"), {"val": color.upper()})
    clr.append(srgb)
    p_pr.append(clr)
    char = p_pr.makeelement(qn("a:buChar"), {"char": "•"})
    p_pr.append(char)


def _cell(
    cell, text: str, size: float, bold: bool, color: str, fill: str, align, font: str, border: str
) -> None:
    from pptx.oxml.ns import qn

    cell.fill.solid()
    cell.fill.fore_color.rgb = _rgb(fill)
    cell.margin_left = cell.margin_right = _pt(8)
    cell.margin_top = cell.margin_bottom = _pt(3)
    frame = cell.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = _pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = _rgb(color)
    # Bordes finos del tema: el estilo por defecto dibuja líneas blancas que
    # en un tema oscuro se ven como una rejilla ajena.
    # El schema exige lnL, lnR, lnT, lnB antes del relleno: fuera de orden,
    # PowerPoint «repara» el archivo.
    tc_pr = cell._tc.get_or_add_tcPr()
    for index, side in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        for node in tc_pr.findall(qn(side)):
            tc_pr.remove(node)
        line = tc_pr.makeelement(qn(side), {"w": str(_pt(0.5)), "cmpd": "sng"})
        solid = line.makeelement(qn("a:solidFill"), {})
        solid.append(solid.makeelement(qn("a:srgbClr"), {"val": border.upper()}))
        line.append(solid)
        tc_pr.insert(index, line)


def _is_number(value: Any) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    text = str(value).strip().replace("\u2212", "-").replace("%", "")
    text = text.replace(",", "").replace(".", "")
    text = text.lstrip("+-$€₡").strip()
    return text.isdigit()


def _decimals(value: float) -> int:
    text = repr(value)
    return min(4, len(text.split(".")[1])) if "." in text and "e" not in text else 0


def _format(value: Any, language: str = "es", decimals: int = 0) -> str:
    """Cifras con los separadores del idioma del documento."""
    if value is None:
        return "\u2014"
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        text = f"{value:,.{decimals if isinstance(value, float) else 0}f}"
        if not language.lower().startswith("en"):
            text = text.replace(",", "\u0000").replace(".", ",").replace("\u0000", ".")
        return text.replace("-", "\u2212")
    return str(value)


def _series_name(item: Any) -> Any:
    return item.get("name") if isinstance(item, dict) else item


def _required(spec: dict[str, Any], key: str) -> Any:
    value = spec.get(key) if isinstance(spec, dict) else None
    if value in (None, ""):
        raise DocumentError(
            DocumentErrorCode.INVALID_SPEC, f"{key} is required", details={"spec": str(spec)[:200]}
        )
    return value


def _strings(spec: dict[str, Any], key: str, low: int, high: int) -> list[str]:
    values = spec.get(key)
    if not isinstance(values, list) or not low <= len(values) <= high:
        raise DocumentError(DocumentErrorCode.INVALID_SPEC, f"{key} needs {low}-{high} items")
    return [str(v) for v in values]


DECK_KEYS = frozenset(
    {"schema_version", "title", "language", "aspect_ratio", "theme", "author", "slides"}
)
COMMON_KEYS = frozenset({"id", "layout", "notes", "source"})
LAYOUT_KEYS: dict[str, frozenset[str]] = {
    "cover": frozenset({"title", "subtitle", "eyebrow", "meta"}),
    "section": frozenset({"title", "subtitle", "number"}),
    "statement": frozenset({"title", "support"}),
    "bullets": frozenset({"title", "points"}),
    "summary": frozenset({"title", "points"}),
    "kpi": frozenset({"title", "kpis", "takeaway"}),
    "chart": frozenset({"title", "chart", "takeaway"}),
    "comparison": frozenset({"title", "columns"}),
    "matrix": frozenset({"title", "quadrants", "x_axis", "y_axis"}),
    "table": frozenset({"title", "columns", "rows", "highlight_row"}),
    "timeline": frozenset({"title", "events"}),
    "process": frozenset({"title", "steps"}),
    "image": frozenset({"title", "image", "caption"}),
    "quote": frozenset({"quote", "author", "role"}),
    "closing": frozenset({"title", "subtitle", "contact"}),
}


def validate_slide(spec: Any, index: int) -> None:
    """Esquema cerrado: un campo desconocido es un error, no se ignora."""
    if not isinstance(spec, dict):
        raise DocumentError(DocumentErrorCode.INVALID_SPEC, f"slide {index} must be an object")
    layout = spec.get("layout", "bullets")
    if layout not in LAYOUTS:
        raise DocumentError(
            DocumentErrorCode.INVALID_SPEC,
            f"slide {index}: unknown layout {layout!r}",
            details={"layouts": list(LAYOUTS)},
        )
    unknown = sorted(set(spec) - COMMON_KEYS - LAYOUT_KEYS[layout])
    if unknown:
        raise DocumentError(
            DocumentErrorCode.INVALID_SPEC,
            f"slide {index} ({layout}): unknown fields {', '.join(unknown)}",
            details={"allowed": sorted(COMMON_KEYS | LAYOUT_KEYS[layout])},
        )


def _blank_layout(prs):
    """El layout con menos marcadores: el en blanco de cada plantilla."""
    layouts = list(prs.slide_layouts)
    for layout in layouts:
        if layout.name.strip().lower() in ("blank", "en blanco"):
            return layout
    return min(layouts, key=lambda layout: len(layout.placeholders))


def add_slide(prs, spec: dict[str, Any], theme_id: str | None, resources: dict[str, str]):
    """Una diapositiva compuesta, añadida al final de un deck existente."""
    builder = _Builder({"theme": theme_id, "slides": [spec]}, resources, prs=prs)
    slide_id = str(spec.get("id") or f"s{len(prs.slides) + 1:02d}")
    slide = builder.add(spec, len(prs.slides) + 1, slide_id)
    return slide, builder.plans[0]


def build(
    spec: dict[str, Any], resources: dict[str, str] | None = None
) -> tuple[bytes, list[SlidePlan]]:
    builder = _Builder(spec, resources or {})
    data = builder.build()
    return data, builder.plans


def plan_findings(plans: list[SlidePlan]) -> list[dict[str, Any]]:
    return [finding for plan in plans for finding in plan.findings]


def template_catalog() -> dict[str, Any]:
    from rinari.documents.design.tokens import THEMES

    return {
        "themes": [t.to_dict() for t in THEMES.values()],
        "layouts": {
            "cover": "title, subtitle?, eyebrow?, meta?",
            "section": "title, subtitle?, number?",
            "statement": "title (one dominant message), support?",
            "bullets": "title, points[1-8]",
            "summary": "title, points[2-6] of text or {head, text}",
            "kpi": "title, kpis[1-4] of {label, value, delta?, trend?, good?, note?}, takeaway?",
            "chart": (
                f"title, chart {{type {'|'.join(CHART_TYPES)}, categories, "
                "series[{name, values}], number_format?}, takeaway?"
            ),
            "comparison": "title, columns[2-3] of {heading, points[1-6]}",
            "matrix": "title, quadrants[4] of {heading, text?, highlight?}, x_axis?, y_axis?",
            "table": "title, columns[≤8], rows[≤14] of cells, highlight_row?",
            "timeline": "title, events[2-6] of {date, heading, text?}",
            "process": "title, steps[2-6] of {heading, text?}",
            "image": "title, image (artifact:// or path), caption?",
            "quote": "quote, author?, role?",
            "closing": "title, subtitle?, contact?",
        },
        "common": "id, layout, notes (speaker notes), source (footnote)",
    }


__all__ = ["LAYOUTS", "Path", "add_slide", "build", "plan_findings", "template_catalog"]
