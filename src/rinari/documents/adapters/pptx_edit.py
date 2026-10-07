"""Edición tipada de presentaciones existentes con python-pptx.

Cada operación apunta a una diapositiva (`slide`, posición desde 1, o
`slide_id`) y, si toca una forma, a la forma (`shape_id` o `shape_name`),
nunca a «el tercer texto que encuentre el parser». Las precondiciones
(`expected_text`, `expected_count`, `expected_title`) se comprueban antes de
tocar nada: si no se cumplen, `REVISION_CONFLICT`.

Cada operación declara las partes OOXML que puede cambiar. El diff de
preservación (preservation.py) compara eso con lo que de verdad cambió.
"""

from __future__ import annotations

import copy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from rinari.documents.contracts import DocumentError, DocumentErrorCode

MAX_OPERATIONS = 200
OPERATIONS = (
    "pptx.set_text",
    "pptx.replace_text",
    "pptx.set_table_cell",
    "pptx.update_chart",
    "pptx.replace_image",
    "pptx.set_notes",
    "pptx.delete_slide",
    "pptx.move_slide",
    "pptx.add_slide",
)
_FIELDS: dict[str, set[str]] = {
    "pptx.set_text": {"slide", "slide_id", "shape_id", "shape_name", "text", "expected_text"},
    "pptx.replace_text": {"find", "replace", "slides", "expected_count", "include_notes"},
    "pptx.set_table_cell": {
        "slide",
        "slide_id",
        "shape_id",
        "shape_name",
        "row",
        "col",
        "text",
        "expected_text",
    },
    "pptx.update_chart": {
        "slide",
        "slide_id",
        "shape_id",
        "shape_name",
        "categories",
        "series",
        "number_format",
        "expected_categories",
    },
    "pptx.replace_image": {"slide", "slide_id", "shape_id", "shape_name", "image", "fit"},
    "pptx.set_notes": {"slide", "slide_id", "text", "expected_text"},
    "pptx.delete_slide": {"slide", "slide_id", "expected_title"},
    "pptx.move_slide": {"slide", "slide_id", "to"},
    "pptx.add_slide": {"after", "spec", "theme"},
}


@dataclass(slots=True)
class EditState:
    """Lo que la edición tocó: partes permitidas y resumen de cambios."""

    allowed: set[str] = field(default_factory=set)
    allow_new: set[str] = field(default_factory=set)
    changes: list[dict[str, Any]] = field(default_factory=list)
    findings: list[dict[str, Any]] = field(default_factory=list)


def _conflict(message: str, **details: Any) -> DocumentError:
    return DocumentError(
        DocumentErrorCode.REVISION_CONFLICT,
        message,
        details=details,
        action="Read the document again and retry with current values",
    )


def _invalid(message: str, **details: Any) -> DocumentError:
    return DocumentError(DocumentErrorCode.INVALID_SPEC, message, details=details)


def _partname(part) -> str:
    return str(part.partname).lstrip("/")


def validate(operations: Any) -> list[dict[str, Any]]:
    """Esquema cerrado de cada operación, antes de abrir el archivo."""
    if not isinstance(operations, list) or not operations:
        raise _invalid("operations must be a non-empty list")
    if len(operations) > MAX_OPERATIONS:
        raise DocumentError(
            DocumentErrorCode.DOCUMENT_LIMIT_EXCEEDED, f"At most {MAX_OPERATIONS} operations"
        )
    for index, op in enumerate(operations, start=1):
        if not isinstance(op, dict):
            raise _invalid(f"operation {index} must be an object")
        name = op.get("op")
        if name not in _FIELDS:
            raise _invalid(f"operation {index}: unknown op {name!r}", ops=list(OPERATIONS))
        unknown = sorted(set(op) - {"op"} - _FIELDS[name])
        if unknown:
            raise _invalid(
                f"operation {index} ({name}): unknown fields {', '.join(unknown)}",
                allowed=sorted(_FIELDS[name]),
            )
    return operations


# -- localización ----------------------------------------------------------------------
def _slide(prs, op: dict[str, Any]):
    slides = list(prs.slides)
    if op.get("slide_id") is not None:
        for slide in slides:
            if slide.slide_id == int(op["slide_id"]):
                return slide
        raise DocumentError(DocumentErrorCode.NOT_FOUND, f"No slide with id {op['slide_id']}")
    position = op.get("slide")
    if not isinstance(position, int) or not 1 <= position <= len(slides):
        raise _invalid(f"slide must be 1..{len(slides)}", slide=position)
    return slides[position - 1]


def _walk(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)


def _shape(slide, op: dict[str, Any]):
    shape_id = op.get("shape_id")
    name = op.get("shape_name")
    if shape_id is None and not name:
        raise _invalid("shape_id or shape_name is required")
    matches = [
        shape
        for shape in _walk(slide.shapes)
        if (shape_id is not None and shape.shape_id == int(shape_id))
        or (shape_id is None and shape.name == name)
    ]
    if not matches:
        raise DocumentError(
            DocumentErrorCode.NOT_FOUND, f"No shape {shape_id or name!r} on that slide"
        )
    if len(matches) > 1:
        raise _invalid(f"shape_name {name!r} is ambiguous; use shape_id")
    return matches[0]


def _position(prs, slide) -> int:
    return [s.slide_id for s in prs.slides].index(slide.slide_id) + 1


# -- texto -------------------------------------------------------------------------------
def _set_paragraph(paragraph, text: str) -> None:
    """Cambia el texto conservando el formato del primer run del párrafo."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run().text = text
        return
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)
    # Saltos y campos sueltos dejarían texto viejo entre medias.
    for node in list(paragraph._p):
        if node.tag.endswith(("}br", "}fld")):
            paragraph._p.remove(node)


def set_frame_text(frame, text: str) -> None:
    lines = str(text).split("\n")
    paragraphs = list(frame.paragraphs)
    template = paragraphs[-1]._p
    for index, line in enumerate(lines):
        if index < len(paragraphs):
            _set_paragraph(paragraphs[index], line)
            continue
        clone = copy.deepcopy(template)
        frame._txBody.append(clone)
        from pptx.text.text import _Paragraph

        _set_paragraph(_Paragraph(clone, frame), line)
    for paragraph in paragraphs[len(lines) :]:
        paragraph._p.getparent().remove(paragraph._p)


def _replace_in_paragraph(paragraph, find: str, replace: str) -> int:
    runs = list(paragraph.runs)
    if not runs:
        return 0
    full = "".join(run.text for run in runs)
    positions = []
    start = full.find(find)
    while start != -1:
        positions.append(start)
        start = full.find(find, start + len(find))
    for start in reversed(positions):
        texts = [run.text for run in runs]
        bounds = []
        cursor = 0
        for text in texts:
            bounds.append(cursor)
            cursor += len(text)
        end = start + len(find)
        first = max(i for i, b in enumerate(bounds) if b <= start and (texts[i] or i == 0))
        last = max(i for i, b in enumerate(bounds) if b < end)
        offset = start - bounds[first]
        if first == last:
            runs[first].text = texts[first][:offset] + replace + texts[first][offset + len(find) :]
            continue
        runs[first].text = texts[first][:offset] + replace
        for middle in range(first + 1, last):
            runs[middle].text = ""
        runs[last].text = texts[last][end - bounds[last] :]
    for run in runs:
        if not run.text and len(paragraph.runs) > 1:
            run._r.getparent().remove(run._r)
    return len(positions)


def _frames(slide, include_notes: bool):
    for shape in _walk(slide.shapes):
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            yield shape.text_frame, slide.part
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame, slide.part
    if include_notes and slide.has_notes_slide:
        yield slide.notes_slide.notes_text_frame, slide.notes_slide.part


# -- operaciones ---------------------------------------------------------------------------
def _op_set_text(prs, op, state: EditState, resources) -> None:
    slide = _slide(prs, op)
    shape = _shape(slide, op)
    if not shape.has_text_frame:
        raise _invalid("That shape has no text; use set_table_cell or update_chart")
    current = shape.text_frame.text
    if "expected_text" in op and current != op["expected_text"]:
        raise _conflict("The shape text changed", current=current[:300])
    set_frame_text(shape.text_frame, str(op.get("text", "")))
    state.allowed.add(_partname(slide.part))
    state.changes.append(
        {
            "op": op["op"],
            "slide": _position(prs, slide),
            "slide_id": slide.slide_id,
            "shape_id": shape.shape_id,
        }
    )


def _op_replace_text(prs, op, state: EditState, resources) -> None:
    find = op.get("find")
    if not isinstance(find, str) or not find:
        raise _invalid("find must be a non-empty string")
    replace = str(op.get("replace", ""))
    wanted = set(op.get("slides") or [])
    total = 0
    touched: list[int] = []
    for position, slide in enumerate(prs.slides, start=1):
        if wanted and position not in wanted:
            continue
        count = 0
        for frame, part in _frames(slide, bool(op.get("include_notes"))):
            hits = sum(_replace_in_paragraph(p, find, replace) for p in frame.paragraphs)
            if hits:
                state.allowed.add(_partname(part))
            count += hits
        if count:
            touched.append(position)
        total += count
    expected = op.get("expected_count")
    if expected is not None and total != int(expected):
        raise _conflict(f"Found {total} occurrences, expected {expected}", found=total)
    if total == 0:
        raise _conflict(f"{find!r} does not appear in the selected slides")
    state.changes.append({"op": op["op"], "count": total, "slides": touched})


def _op_set_table_cell(prs, op, state: EditState, resources) -> None:
    slide = _slide(prs, op)
    shape = _shape(slide, op)
    if not getattr(shape, "has_table", False) or not shape.has_table:
        raise _invalid("That shape is not a table")
    table = shape.table
    row, col = op.get("row"), op.get("col")
    if not isinstance(row, int) or not isinstance(col, int):
        raise _invalid("row and col are 0-based integers")
    if not (0 <= row < len(table.rows) and 0 <= col < len(table.columns)):
        raise _invalid(f"Cell out of range ({len(table.rows)}x{len(table.columns)})")
    cell = table.cell(row, col)
    if "expected_text" in op and cell.text != op["expected_text"]:
        raise _conflict("The cell changed", current=cell.text[:300])
    set_frame_text(cell.text_frame, str(op.get("text", "")))
    state.allowed.add(_partname(slide.part))
    state.changes.append(
        {
            "op": op["op"],
            "slide": _position(prs, slide),
            "shape_id": shape.shape_id,
            "cell": [row, col],
        }
    )


def _op_update_chart(prs, op, state: EditState, resources) -> None:
    from pptx.chart.data import CategoryChartData

    slide = _slide(prs, op)
    shape = _shape(slide, op)
    if not getattr(shape, "has_chart", False) or not shape.has_chart:
        raise _invalid("That shape is not a chart")
    chart = shape.chart
    plot = chart.plots[0] if chart.plots else None
    if plot is None:
        raise DocumentError(DocumentErrorCode.UNSUPPORTED_FEATURE, "Chart without a plot")
    from pptx.enum.chart import XL_CHART_TYPE

    if (
        chart.chart_type
        in (
            XL_CHART_TYPE.XY_SCATTER,
            XL_CHART_TYPE.BUBBLE,
        )
        or "scatter" in str(chart.chart_type).lower()
    ):
        raise DocumentError(
            DocumentErrorCode.UNSUPPORTED_FEATURE, "XY/bubble charts are not editable yet"
        )
    categories = op.get("categories")
    series = op.get("series")
    if not isinstance(categories, list) or not categories:
        raise _invalid("categories must be a non-empty list")
    if not isinstance(series, list) or not series:
        raise _invalid("series must be a non-empty list of {name, values}")
    current = [str(c) for c in plot.categories]
    if "expected_categories" in op and [str(c) for c in op["expected_categories"]] != current:
        raise _conflict("The chart categories changed", current=current)
    data = CategoryChartData(number_format=op.get("number_format") or _number_format(chart))
    data.categories = categories
    for item in series:
        if not isinstance(item, dict) or "values" not in item:
            raise _invalid("each series needs name and values")
        values = item["values"]
        if not isinstance(values, list) or len(values) != len(categories):
            raise _invalid(f"series {item.get('name')!r} must have {len(categories)} values")
        data.add_series(str(item.get("name") or ""), values)
    chart.replace_data(data)
    chart_part = shape.chart_part
    state.allowed.add(_partname(chart_part))
    workbook = chart_part.chart_workbook.xlsx_part
    if workbook is not None:
        state.allowed.add(_partname(workbook))
    state.allow_new.add("ppt/embeddings/")
    state.changes.append(
        {
            "op": op["op"],
            "slide": _position(prs, slide),
            "shape_id": shape.shape_id,
            "series": len(series),
            "categories": len(categories),
        }
    )


def _number_format(chart) -> str:
    """El formato de las cifras que ya tenía el gráfico (p. ej. 0"%")."""
    from pptx.oxml.ns import qn

    for code in chart._chartSpace.iter(qn("c:formatCode")):
        if code.text:
            return code.text
    return "General"


def _op_replace_image(prs, op, state: EditState, resources) -> None:
    from PIL import Image
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.oxml.ns import qn

    slide = _slide(prs, op)
    shape = _shape(slide, op)
    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
        raise _invalid("That shape is not a picture")
    path = resources.get(str(op.get("image")))
    if path is None:
        raise DocumentError(DocumentErrorCode.NOT_FOUND, f"Image {op.get('image')!r} not provided")
    blip = shape._element.blipFill.find(qn("a:blip"))
    old = blip.get(qn("r:embed"))
    old_part = slide.part.related_part(old) if old else None
    _, rid = slide.part.get_or_add_image_part(path)
    blip.set(qn("r:embed"), rid)
    # Llenar la caja sin deformar: recorta el sobrante (cover) o lo deja entero.
    with Image.open(path) as image:
        iw, ih = image.size
    box = shape.width / shape.height
    ratio = iw / ih
    shape.crop_left = shape.crop_right = shape.crop_top = shape.crop_bottom = 0.0
    if op.get("fit", "cover") == "cover":
        if ratio > box:
            excess = (1 - box / ratio) / 2
            shape.crop_left = shape.crop_right = excess
        elif ratio < box:
            excess = (1 - ratio / box) / 2
            shape.crop_top = shape.crop_bottom = excess
    # La imagen vieja se suelta solo si ya nada de la diapositiva la usa.
    if old and old != rid and slide.part._rel_ref_count(old) == 0:
        slide.part._rels.pop(old)
    state.allowed.add(_partname(slide.part))
    if old_part is not None and old not in slide.part.rels:
        state.allowed.add(_partname(old_part))
    state.allow_new.add("ppt/media/")
    state.changes.append(
        {"op": op["op"], "slide": _position(prs, slide), "shape_id": shape.shape_id}
    )


def _op_set_notes(prs, op, state: EditState, resources) -> None:
    slide = _slide(prs, op)
    had_master = prs.part.package is not None and any(
        str(p.partname).startswith("/ppt/notesMasters/") for p in prs.part.package.iter_parts()
    )
    current = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
    if "expected_text" in op and current != op["expected_text"]:
        raise _conflict("The notes changed", current=current[:300])
    frame = slide.notes_slide.notes_text_frame
    set_frame_text(frame, str(op.get("text", "")))
    state.allowed.add(_partname(slide.notes_slide.part))
    state.allowed.add(_partname(slide.part))
    state.allow_new.add("ppt/notesSlides/")
    if not had_master:
        state.allowed.add("ppt/presentation.xml")
        state.allow_new.update({"ppt/notesMasters/", "ppt/theme/"})
    state.changes.append({"op": op["op"], "slide": _position(prs, slide)})


def _reachable(part, stop: set[str]) -> set[str]:
    """Partes que cuelgan de una diapositiva (sin subir a layouts ni masters)."""
    seen: set[str] = set()
    pending = [part]
    while pending:
        current = pending.pop()
        name = _partname(current)
        if name in seen:
            continue
        seen.add(name)
        for rel in current.rels.values():
            if rel.is_external:
                continue
            target = rel.target_part
            if any(_partname(target).startswith(prefix) for prefix in stop):
                continue
            pending.append(target)
    return seen


def _op_delete_slide(prs, op, state: EditState, resources) -> None:
    slide = _slide(prs, op)
    if len(prs.slides) < 2:
        raise _invalid("A presentation keeps at least one slide")
    if "expected_title" in op:
        from rinari.documents.adapters.pptx_read import _title

        if _title(slide) != str(op["expected_title"]).strip():
            raise _conflict("The slide title changed", current=_title(slide))
    position = _position(prs, slide)
    slide_id = slide.slide_id
    stop = {"ppt/slideLayouts/", "ppt/slideMasters/", "ppt/notesMasters/", "ppt/theme/"}
    state.allowed |= _reachable(slide.part, stop)
    sld_ids = prs.slides._sldIdLst
    for sld in list(sld_ids):
        if int(sld.get("id")) == slide_id:
            rid = sld.rId
            sld_ids.remove(sld)
            prs.part.drop_rel(rid)
            break
    state.allowed.add("ppt/presentation.xml")
    state.changes.append({"op": op["op"], "slide": position, "slide_id": slide_id})


def _op_move_slide(prs, op, state: EditState, resources) -> None:
    slide = _slide(prs, op)
    to = op.get("to")
    count = len(prs.slides)
    if not isinstance(to, int) or not 1 <= to <= count:
        raise _invalid(f"to must be 1..{count}")
    sld_ids = prs.slides._sldIdLst
    element = next(s for s in sld_ids if int(s.get("id")) == slide.slide_id)
    origin = list(sld_ids).index(element) + 1
    sld_ids.remove(element)
    sld_ids.insert(to - 1, element)
    state.allowed.add("ppt/presentation.xml")
    state.changes.append({"op": op["op"], "slide_id": slide.slide_id, "from": origin, "to": to})


def _op_add_slide(prs, op, state: EditState, resources) -> None:
    from rinari.documents.adapters import pptx_build

    spec = op.get("spec")
    if not isinstance(spec, dict):
        raise _invalid("spec must be a slide object (same as a DeckSpec slide)")
    after = op.get("after", len(prs.slides))
    if not isinstance(after, int) or not 0 <= after <= len(prs.slides):
        raise _invalid(f"after must be 0..{len(prs.slides)}")
    slide, plan = pptx_build.add_slide(prs, spec, op.get("theme"), resources)
    sld_ids = prs.slides._sldIdLst
    element = sld_ids[-1]
    sld_ids.remove(element)
    sld_ids.insert(after, element)
    state.allowed.add("ppt/presentation.xml")
    state.allow_new.update({"ppt/slides/", "ppt/charts/", "ppt/embeddings/", "ppt/media/"})
    state.allow_new.update({"ppt/notesSlides/", "ppt/notesMasters/", "ppt/theme/"})
    state.findings.extend(plan.findings)
    state.changes.append(
        {
            "op": op["op"],
            "slide": after + 1,
            "slide_id": slide.slide_id,
            "layout": plan.layout,
        }
    )


_HANDLERS = {
    "pptx.set_text": _op_set_text,
    "pptx.replace_text": _op_replace_text,
    "pptx.set_table_cell": _op_set_table_cell,
    "pptx.update_chart": _op_update_chart,
    "pptx.replace_image": _op_replace_image,
    "pptx.set_notes": _op_set_notes,
    "pptx.delete_slide": _op_delete_slide,
    "pptx.move_slide": _op_move_slide,
    "pptx.add_slide": _op_add_slide,
}


def apply(
    source: Path, target: Path, operations: list[dict[str, Any]], resources: dict[str, str]
) -> EditState:
    """Aplica todas las operaciones o ninguna: la salida es otro archivo."""
    from pptx import Presentation

    validate(operations)
    prs = Presentation(str(source))
    state = EditState()
    for index, op in enumerate(operations, start=1):
        try:
            _HANDLERS[op["op"]](prs, op, state, resources)
        except DocumentError as exc:
            exc.details = {**(exc.details or {}), "operation": index, "op": op["op"]}
            raise
    prs.save(str(target))
    return state


__all__ = ["OPERATIONS", "EditState", "apply", "set_frame_text", "validate"]
