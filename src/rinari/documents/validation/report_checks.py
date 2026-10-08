"""Comprobaciones de documentos Word y PDF.

Estructura (abre, relaciones), contenido esperado, campos (un índice escrito
no es un índice paginado) y, en PDF, capa de texto, fuentes incrustadas y
formularios (valor guardado y apariencia dibujada).
"""

from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any

from rinari.documents.contracts import (
    CHECK_FAILED,
    CHECK_NOT_APPLICABLE,
    CHECK_NOT_RUN,
    CHECK_PARTIAL,
    CHECK_PASSED,
    Check,
)
from rinari.documents.validation.pptx_checks import relationship_findings

_STANDARD_FONTS = {
    "Courier",
    "Courier-Bold",
    "Courier-Oblique",
    "Courier-BoldOblique",
    "Helvetica",
    "Helvetica-Bold",
    "Helvetica-Oblique",
    "Helvetica-BoldOblique",
    "Times-Roman",
    "Times-Bold",
    "Times-Italic",
    "Times-BoldItalic",
    "Symbol",
    "ZapfDingbats",
}


# -- DOCX ----------------------------------------------------------------------------------
def _docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    for section in document.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(parts)


def docx_structure(path: Path) -> Check:
    try:
        with zipfile.ZipFile(path) as archive:
            findings = relationship_findings(archive)
        from docx import Document

        document = Document(str(path))
        blocks = len(document.element.body)
    except Exception as exc:
        return Check(CHECK_FAILED, reason="UNREADABLE", findings=[{"detail": str(exc)[:300]}])
    return Check(
        CHECK_FAILED if findings else CHECK_PASSED,
        evidence={"blocks": blocks},
        findings=findings[:50],
    )


def docx_fields(path: Path, *, updated_by: str | None) -> Check:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8", "replace")
    tocs = xml.count(" TOC ")
    if not tocs:
        return Check(CHECK_NOT_APPLICABLE, evidence={"toc": 0})
    if updated_by:
        return Check(CHECK_PASSED, evidence={"toc": tocs, "updated_by": updated_by})
    return Check(CHECK_NOT_RUN, reason="FIELDS_NOT_UPDATED", evidence={"toc": tocs})


def text_content(text: str, expected: list[str] | None) -> Check:
    if not expected:
        return Check(CHECK_NOT_RUN, reason="NO_CRITERIA")
    folded = text.casefold()
    missing = [item for item in expected if str(item).casefold() not in folded]
    return Check(
        CHECK_FAILED if missing else CHECK_PASSED,
        evidence={"expected": len(expected)},
        findings=[{"code": "MISSING_TEXT", "severity": "error", "text": m} for m in missing],
    )


def docx_checks(path: Path, expected: list[str] | None, *, fields_updated_by: str | None = None):
    return {
        "structure": docx_structure(path).to_dict(),
        "content": text_content(_docx_text(path), expected).to_dict(),
        "fields": docx_fields(path, updated_by=fields_updated_by).to_dict(),
    }


def docx_blocks(path: Path) -> list[str]:
    from docx import Document

    from rinari.documents.adapters.basic_read import _blocks

    out = []
    for block in _blocks(Document(str(path))):
        if hasattr(block, "rows"):
            out.append("\t".join(cell.text for row in block.rows for cell in row.cells)[:2000])
        else:
            out.append(block.text[:2000])
    return out


def docx_semantic_diff(before: Path, after: Path) -> list[dict[str, Any]]:
    import difflib

    a, b = docx_blocks(before), docx_blocks(after)
    out: list[dict[str, Any]] = []
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(a=a, b=b, autojunk=False).get_opcodes():
        if tag == "equal":
            continue
        if tag == "replace" and i2 - i1 == j2 - j1:
            for offset in range(i2 - i1):
                out.append(
                    {
                        "change": "content_changed",
                        "block": j1 + offset + 1,
                        "before": a[i1 + offset][:300],
                        "after": b[j1 + offset][:300],
                    }
                )
            continue
        for index in range(i1, i2):
            out.append({"change": "block_removed", "block": index + 1, "before": a[index][:300]})
        for index in range(j1, j2):
            out.append({"change": "block_added", "block": index + 1, "after": b[index][:300]})
    return out[:300]


# -- PDF ------------------------------------------------------------------------------------
def pdf_structure(path: Path) -> Check:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path), strict=False)
        if reader.is_encrypted:
            return Check(CHECK_FAILED, reason="ENCRYPTED")
        pages = len(reader.pages)
    except Exception as exc:
        return Check(CHECK_FAILED, reason="UNREADABLE", findings=[{"detail": str(exc)[:300]}])
    return Check(CHECK_PASSED if pages else CHECK_FAILED, evidence={"pages": pages})


def pdf_text_and_fonts(path: Path) -> Check:
    """Texto seleccionable en cada página y fuentes incrustadas."""
    from pypdf import PdfReader

    reader = PdfReader(str(path), strict=False)
    findings: list[dict[str, Any]] = []
    without_text = []
    fonts: dict[str, bool] = {}
    for number, page in enumerate(reader.pages[:2_000], start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            without_text.append(number)
        resources = page.get("/Resources") or {}
        for font in (resources.get("/Font") or {}).values() if resources else []:
            font = font.get_object()
            name = str(font.get("/BaseFont", "")).lstrip("/").split("+")[-1]
            descriptor = font.get("/FontDescriptor")
            if descriptor is None and font.get("/DescendantFonts"):
                descriptor = font["/DescendantFonts"][0].get_object().get("/FontDescriptor")
            embedded = False
            if descriptor is not None:
                descriptor = descriptor.get_object()
                embedded = any(k in descriptor for k in ("/FontFile", "/FontFile2", "/FontFile3"))
            fonts[name] = fonts.get(name, False) or embedded or name in _STANDARD_FONTS
    for name, ok in fonts.items():
        if not ok:
            findings.append({"code": "FONT_NOT_EMBEDDED", "severity": "warning", "font": name})
    if without_text:
        findings.append(
            {
                "code": "NO_TEXT_LAYER",
                "severity": "warning",
                "pages": without_text[:50],
                "count": len(without_text),
            }
        )
    status = CHECK_PASSED if not findings else CHECK_PARTIAL
    return Check(
        status,
        evidence={"pages": len(reader.pages), "fonts": sorted(fonts)[:30]},
        findings=findings,
    )


def pdf_forms(path: Path, filled: dict[str, str] | None) -> Check:
    if not filled:
        return Check(CHECK_NOT_APPLICABLE)
    from rinari.documents.adapters.pdf_edit import verify_form

    result = verify_form(path, filled)
    if result["stored"] and result["visible"]:
        return Check(
            CHECK_PASSED, evidence={"fields": len(filled), "stored": True, "visible": True}
        )
    findings = [
        {"code": "FIELD_NOT_STORED", "severity": "error", "field": n} for n in result["wrong"]
    ]
    findings += [
        {"code": "FIELD_NOT_VISIBLE", "severity": "error", "field": n} for n in result["invisible"]
    ]
    return Check(CHECK_FAILED, evidence={"fields": len(filled)}, findings=findings)


def pdf_checks(path: Path, expected: list[str] | None, *, filled: dict[str, str] | None = None):
    from rinari.documents.adapters.pdf_edit import page_texts

    text = "\n".join(page_texts(path, limit=20_000))
    return {
        "structure": pdf_structure(path).to_dict(),
        "text_layer": pdf_text_and_fonts(path).to_dict(),
        "content": text_content(text, expected).to_dict(),
        "forms": pdf_forms(path, filled).to_dict(),
    }


def pdf_semantic_diff(before: Path, after: Path) -> list[dict[str, Any]]:
    import difflib

    from rinari.documents.adapters.pdf_edit import page_texts

    a, b = page_texts(before, 500), page_texts(after, 500)
    out: list[dict[str, Any]] = []
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(a=a, b=b, autojunk=False).get_opcodes():
        if tag == "equal":
            continue
        for index in range(i1, i2):
            out.append({"change": "page_removed", "page": index + 1})
        for index in range(j1, j2):
            out.append({"change": "page_added", "page": index + 1})
    if len(a) != len(b):
        out.append({"change": "page_count", "before": len(a), "after": len(b)})
    return out[:300]


__all__ = [
    "docx_checks",
    "docx_semantic_diff",
    "pdf_checks",
    "pdf_semantic_diff",
]
