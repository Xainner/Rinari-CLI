"""Comprobaciones de un libro: estructura, contenido y fórmulas con su caché.

Una fórmula escrita no es una fórmula calculada: `formulas` solo pasa cuando
el backend certificado (Excel) recalculó esta revisión y ninguna celda quedó
en error. Sin ese cálculo, el estado es «sin ejecutar», con el motivo.
"""

from __future__ import annotations

import warnings
import zipfile
from pathlib import Path
from typing import Any

from rinari.documents.contracts import (
    CHECK_FAILED,
    CHECK_NOT_APPLICABLE,
    CHECK_NOT_RUN,
    CHECK_PASSED,
    Check,
)
from rinari.documents.validation.pptx_checks import relationship_findings

CONTENT_CELLS = 200_000


def structure(path: Path) -> Check:
    import openpyxl

    warnings.filterwarnings("ignore", module="openpyxl")
    try:
        with zipfile.ZipFile(path) as archive:
            findings = relationship_findings(archive)
        workbook = openpyxl.load_workbook(str(path), read_only=True)
        sheets = len(workbook.sheetnames)
        workbook.close()
    except Exception as exc:
        return Check(CHECK_FAILED, reason="UNREADABLE", findings=[{"detail": str(exc)[:300]}])
    return Check(
        CHECK_FAILED if findings else CHECK_PASSED,
        evidence={"sheets": sheets},
        findings=findings[:50],
    )


def content(path: Path, expected: list[str] | None) -> Check:
    if not expected:
        return Check(CHECK_NOT_RUN, reason="NO_CRITERIA")
    import openpyxl

    warnings.filterwarnings("ignore", module="openpyxl")
    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=False)
    seen: list[str] = []
    scanned = 0
    try:
        for sheet in workbook.worksheets:
            seen.append(sheet.title)
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    if value is not None:
                        seen.append(str(value))
                scanned += len(row)
                if scanned > CONTENT_CELLS:
                    break
    finally:
        workbook.close()
    text = "\n".join(seen).casefold()
    missing = [item for item in expected if str(item).casefold() not in text]
    return Check(
        CHECK_FAILED if missing else CHECK_PASSED,
        evidence={"expected": len(expected), "cells_scanned": scanned},
        findings=[{"code": "MISSING_TEXT", "severity": "error", "text": m} for m in missing],
    )


def formulas_check(inventory: dict[str, Any], *, calculated_by: str | None) -> Check:
    """Del inventario de fórmulas y de quién guardó la revisión, el estado honesto."""
    evidence = {
        k: inventory.get(k) for k in ("formulas", "cached", "missing", "cache_state", "truncated")
    }
    if not inventory.get("formulas"):
        return Check(CHECK_NOT_APPLICABLE, evidence=evidence)
    errors = inventory.get("errors") or []
    if errors:
        return Check(
            CHECK_FAILED,
            reason="FORMULA_ERRORS",
            evidence={**evidence, "backend": calculated_by},
            findings=[
                {
                    "code": "FORMULA_ERROR",
                    "severity": "error",
                    "cell": e["cell"],
                    "message": e["error"],
                }
                for e in errors
            ],
        )
    if calculated_by and inventory.get("cache_state") == "fresh":
        return Check(CHECK_PASSED, evidence={**evidence, "backend": calculated_by})
    return Check(
        CHECK_NOT_RUN,
        reason="NOT_CALCULATED",
        evidence=evidence,
    )


__all__ = ["content", "formulas_check", "structure"]
