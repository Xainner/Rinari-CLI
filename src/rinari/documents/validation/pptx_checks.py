"""Comprobaciones de una presentación, generada por Rinari o ajena.

Estructura (abre, relaciones resueltas), composición (límites, solapes no
previstos, tamaños mínimos, desbordes estimados, contraste) y datos de
gráficos. La revisión visual de verdad es otra dimensión: la registra
`documents.review` sobre los renders de esta revisión exacta.
"""

from __future__ import annotations

import posixpath
import re
import zipfile
from pathlib import Path
from typing import Any

from rinari.documents.contracts import (
    CHECK_FAILED,
    CHECK_NOT_RUN,
    CHECK_PASSED,
    Check,
)

EMU_PER_PT = 12700
MIN_FONT_PT = 10
TOLERANCE_PT = 1.5


def structure(path: Path) -> Check:
    findings: list[dict[str, Any]] = []
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            from defusedxml import ElementTree

            for rels in [n for n in names if n.endswith(".rels")]:
                base = posixpath.dirname(posixpath.dirname(rels))
                root = ElementTree.fromstring(archive.read(rels))
                for node in root:
                    if node.attrib.get("TargetMode") == "External":
                        continue
                    target = node.attrib.get("Target", "")
                    resolved = (
                        target.lstrip("/")
                        if target.startswith("/")
                        else posixpath.normpath(posixpath.join(base, target))
                    )
                    if resolved not in names:
                        findings.append(
                            {
                                "code": "BROKEN_RELATIONSHIP",
                                "severity": "error",
                                "part": rels,
                                "target": target,
                            }
                        )
        from pptx import Presentation

        presentation = Presentation(str(path))
        slides = len(presentation.slides)
    except Exception as exc:
        return Check(CHECK_FAILED, reason="UNREADABLE", findings=[{"detail": str(exc)[:300]}])
    status = CHECK_FAILED if findings else CHECK_PASSED
    return Check(status, evidence={"slides": slides}, findings=findings[:50])


def _background(slide) -> str:
    try:
        fill = slide.background.fill
        if fill.type == 1:  # MSO_FILL.SOLID
            return str(fill.fore_color.rgb)
    except Exception:
        pass
    return "FFFFFF"


def _box(shape) -> tuple[float, float, float, float] | None:
    if shape.left is None or shape.top is None or shape.width is None or shape.height is None:
        return None
    return (
        shape.left / EMU_PER_PT,
        shape.top / EMU_PER_PT,
        shape.width / EMU_PER_PT,
        shape.height / EMU_PER_PT,
    )


def _overlap(a, b) -> float:
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    w = min(ax + aw, bx + bw) - max(ax, bx)
    h = min(ay + ah, by + bh) - max(ay, by)
    return max(0.0, w) * max(0.0, h)


def _shape_fill(shape) -> str | None:
    try:
        if shape.fill.type == 1:
            return str(shape.fill.fore_color.rgb)
    except Exception:
        return None
    return None


def layout(path: Path) -> Check:
    """Composición estática de todas las diapositivas."""
    from pptx import Presentation

    from rinari.documents.design import measure
    from rinari.documents.design.tokens import contrast

    presentation = Presentation(str(path))
    width = presentation.slide_width / EMU_PER_PT
    height = presentation.slide_height / EMU_PER_PT
    findings: list[dict[str, Any]] = []
    for index, slide in enumerate(presentation.slides, start=1):
        background = _background(slide)
        texts: list[tuple[Any, tuple[float, float, float, float]]] = []
        fills: list[tuple[tuple[float, float, float, float], str]] = []
        for shape in slide.shapes:
            box = _box(shape)
            if box is None:
                continue
            x, y, w, h = box
            if (
                x < -TOLERANCE_PT
                or y < -TOLERANCE_PT
                or x + w > width + TOLERANCE_PT
                or y + h > height + TOLERANCE_PT
            ) and shape.rotation == 0:
                findings.append(_finding("OUT_OF_BOUNDS", "error", index, shape))
            fill = _shape_fill(shape)
            if fill:
                fills.append((box, fill))
            if getattr(shape, "has_chart", False) and shape.has_chart:
                findings.extend(_chart_findings(index, shape))
            if not getattr(shape, "has_text_frame", False) or not shape.text_frame.text.strip():
                continue
            if shape.rotation == 0 and not _decorative(shape):
                texts.append((shape, box))
            sizes = [
                run.font.size.pt
                for paragraph in shape.text_frame.paragraphs
                for run in paragraph.runs
                if run.font.size is not None
            ]
            if sizes and min(sizes) < MIN_FONT_PT:
                findings.append(
                    _finding("TEXT_TOO_SMALL", "warning", index, shape, size=min(sizes))
                )
            size = max(sizes) if sizes else 18.0
            family = _font(shape) or "Calibri"
            needed, _ = measure.text_height(shape.text_frame.text, w - 10, family, size)
            if needed > h * 1.15 + 6 and not _autofit(shape):
                findings.append(
                    _finding(
                        "TEXT_OVERFLOW",
                        "error",
                        index,
                        shape,
                        needed_pt=round(needed, 1),
                        box_pt=round(h, 1),
                    )
                )
            color = _text_color(shape)
            if color:
                behind = _behind(box, fills) or background
                ratio = contrast(color, behind)
                large = size >= 24 or (size >= 18.5 and _bold(shape))
                if ratio < (3.0 if large else 4.5):
                    findings.append(
                        _finding("LOW_CONTRAST", "error", index, shape, ratio=round(ratio, 2))
                    )
        for i, (a, box_a) in enumerate(texts):
            for b, box_b in texts[i + 1 :]:
                area = _overlap(box_a, box_b)
                smaller = min(box_a[2] * box_a[3], box_b[2] * box_b[3]) or 1
                if area / smaller > 0.15:
                    findings.append(
                        {
                            "code": "TEXT_OVERLAP",
                            "severity": "warning",
                            "slide": index,
                            "shapes": [a.shape_id, b.shape_id],
                            "names": [a.name, b.name],
                        }
                    )
    errors = [f for f in findings if f["severity"] == "error"]
    status = CHECK_FAILED if errors else CHECK_PASSED
    return Check(
        status,
        evidence={"slides": len(presentation.slides), "issues": len(findings)},
        findings=findings[:200],
    )


def _finding(code: str, severity: str, slide: int, shape, **extra: Any) -> dict[str, Any]:
    return {
        "code": code,
        "severity": severity,
        "slide": slide,
        "shape_id": shape.shape_id,
        "name": shape.name,
        **extra,
    }


def _decorative(shape) -> bool:
    """Un adorno declarado por el layout (comillas grandes): su solape es intencional."""
    return bool(_DECORATIVE.match(shape.name or ""))


_DECORATIVE = re.compile(r"^rinari:[^:]+:(mark|deco-[\w-]+)$")


def _font(shape) -> str | None:
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.name:
                return run.font.name
    return None


def _bold(shape) -> bool:
    runs = [run for p in shape.text_frame.paragraphs for run in p.runs]
    return bool(runs) and all(run.font.bold for run in runs)


def _text_color(shape) -> str | None:
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            try:
                if run.font.color and run.font.color.type is not None and run.font.color.rgb:
                    return str(run.font.color.rgb)
            except Exception:
                continue
    return None


def _autofit(shape) -> bool:
    body = shape.text_frame._txBody.bodyPr
    return any(child.tag.endswith(("normAutofit", "spAutoFit")) for child in body)


def _behind(box, fills) -> str | None:
    x, y, w, h = box
    cx, cy = x + w / 2, y + h / 2
    found = None
    for (fx, fy, fw, fh), color in fills:
        if fx <= cx <= fx + fw and fy <= cy <= fy + fh and fw * fh >= w * h * 0.9:
            found = color
    return found


def _chart_findings(slide: int, shape) -> list[dict[str, Any]]:
    out = []
    try:
        plot = shape.chart.plots[0]
        categories = len(list(plot.categories))
        for series in plot.series:
            values = list(series.values)
            if len(values) != categories:
                out.append(
                    _finding(
                        "CHART_SERIES_MISMATCH",
                        "error",
                        slide,
                        shape,
                        series=series.name,
                        values=len(values),
                        categories=categories,
                    )
                )
            if values and all(v is None for v in values):
                out.append(
                    _finding("CHART_EMPTY_SERIES", "error", slide, shape, series=series.name)
                )
    except Exception as exc:
        out.append(_finding("CHART_UNREADABLE", "error", slide, shape, detail=str(exc)[:200]))
    return out


def content(path: Path, expected: list[str] | None) -> Check:
    if not expected:
        return Check(CHECK_NOT_RUN, reason="NO_CRITERIA")
    from pptx import Presentation

    presentation = Presentation(str(path))
    corpus = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                corpus.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                corpus.extend(cell.text for row in shape.table.rows for cell in row.cells)
        if slide.has_notes_slide:
            corpus.append(slide.notes_slide.notes_text_frame.text)
    text = "\n".join(corpus).casefold()
    missing = [item for item in expected if str(item).casefold() not in text]
    return Check(
        CHECK_FAILED if missing else CHECK_PASSED,
        evidence={"expected": len(expected)},
        findings=[{"code": "MISSING_TEXT", "severity": "error", "text": m} for m in missing],
    )
